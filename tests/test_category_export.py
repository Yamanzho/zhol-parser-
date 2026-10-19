"""Category and search query stay separate in the platform workbook.

These tests use the isolated database from conftest. They do not open a map site
and they do not write category text back onto a snapshot.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.platform_rows import list_platform_rows, write_platform_workbook
from zhol.browser.maps.visible import parse_card
from zhol.db.models import CardSnapshot, CollectionBatch, Company
from zhol.timeutil import now

ROOT = Path(__file__).resolve().parents[1]
ONE = "tests/fixtures/category-recovery/one.html"
MANY = "tests/fixtures/category-recovery/many.html"


@pytest.fixture
def factory(migrated_engine):
    return sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)


def _company(session, name: str) -> Company:
    company = Company(name=name, city="Алматы", niche="не использовать как категорию", source="saved_pages")
    session.add(company)
    session.flush()
    return company


def _card(session, company: Company, platform: str, fields: dict, sources: dict, url: str) -> CardSnapshot:
    row = CardSnapshot(
        company_id=company.id,
        platform=platform,
        match_status="found",
        url=url,
        fields=fields,
        field_sources=sources,
        purpose="prospect",
        collected_at=now(),
        collected_by="test",
        status="collected",
    )
    session.add(row)
    session.flush()
    return row


def test_a_single_source_label_is_kept_even_outside_the_dictionary() -> None:
    fields, notes = parse_card("Бар на Абая\nБар\n4.8\n", name="Бар на Абая", hrefs=[], platform="google")
    assert fields.get("category") == "Бар"
    assert notes.get("category") == "shown"


def test_workbook_keeps_category_query_and_does_not_guess(factory, tmp_path: Path) -> None:
    with factory() as session, session.begin():
        named = _company(session, "С категорией")
        blank = _company(session, "Без категории")
        recoverable = _company(session, "Одна рубрика в файле")
        mixed = _company(session, "Несколько рубрик")
        batch = CollectionBatch(
            city="Алматы",
            category="барбершоп",
            platforms=["2gis"],
            item_limit=5,
            source_mode="browser_2gis",
            status="completed",
            out_path="/tmp/zhol-category-test.xlsx",
            catalog_path="browser_2gis",
        )
        session.add(batch)
        session.flush()
        stored = _card(
            session,
            named,
            "2gis",
            {"name": "С категорией", "category": "Парикмахерские услуги", "city": "Алматы", "phones": ["+77010000001"]},
            {"_collection": {"key": f"browser_2gis:{batch.id}:1", "search_query": "салон красоты"}},
            "https://2gis.kz/almaty/firm/1",
        )
        _card(
            session,
            blank,
            "2gis",
            {"name": "Без категории", "city": "Алматы"},
            {"_collection": {"key": f"batch:{batch.id}:2"}},
            "https://2gis.kz/almaty/firm/2",
        )
        _card(
            session,
            named,
            "google",
            {"name": "Google Name", "category": "Beauty salon", "city": "Алматы"},
            {"_collection": {"search_query": "beauty salon"}},
            "https://www.google.com/maps/place/1",
        )
        saved = _card(
            session,
            recoverable,
            "yandex",
            {"name": "Одна рубрика в файле"},
            {"category": {"reason": "unknown_label"}, "_collection": {"raw_path": ONE}},
            "https://yandex.kz/maps/org/1",
        )
        crowded = _card(
            session,
            mixed,
            "yandex",
            {"name": "Несколько рубрик"},
            {"category": {"reason": "unknown_label"}, "_collection": {"raw_path": MANY}},
            "https://yandex.kz/maps/org/2",
        )
        before = (dict(stored.fields), dict(saved.field_sources), dict(crowded.field_sources))
        gis = list_platform_rows(session, "2gis")
        google = list_platform_rows(session, "google")
        path = tmp_path / "platforms.xlsx"
        counts = write_platform_workbook(session, path)
        session.refresh(stored)
        session.refresh(saved)
        session.refresh(crowded)
        assert (dict(stored.fields), dict(saved.field_sources), dict(crowded.field_sources)) == before

    by_name = {row["name"]: row for row in gis["items"]}
    assert by_name["С категорией"]["category"] == "Парикмахерские услуги"
    assert by_name["С категорией"]["search_query"] == "салон красоты"
    assert by_name["Без категории"]["category"] == "Не указана"
    assert by_name["Без категории"]["search_query"] == "барбершоп"
    assert "Google Name" not in by_name
    assert google["items"][0]["category"] == "Beauty salon"
    assert google["items"][0]["search_query"] == "beauty salon"
    assert counts == {"2gis": 2, "google": 1}

    book = load_workbook(path)
    assert book.sheetnames == ["2ГИС", "Google Maps"]
    headers = [cell.value for cell in book["2ГИС"][1]]
    assert headers == [
        "Название",
        "Категория",
        "Поисковый запрос",
        "Источник",
        "Город",
        "Регион",
        "Адрес",
        "Телефон",
        "WhatsApp",
        "Instagram",
        "Ссылка на карточку",
        "Рейтинг",
        "Отзывы",
        "Дата сбора",
        "Полнота",
    ]
    rows = {
        row[0]: row
        for row in book["2ГИС"].iter_rows(min_row=2, values_only=True)
    }
    assert rows["С категорией"][1] == "Парикмахерские услуги"
    assert rows["С категорией"][2] == "салон красоты"
    assert rows["С категорией"][3] == "2ГИС"
    assert rows["С категорией"][4] == "Алматы"
    assert rows["С категорией"][5] == "город Алматы"
    assert rows["С категорией"][7] == "'+77010000001"
    assert rows["Без категории"][1] == "Не указана"
    assert rows["Без категории"][2] == "барбершоп"
    google_names = [row[0] for row in book["Google Maps"].iter_rows(min_row=2, values_only=True)]
    assert google_names == ["Google Name"]


def test_saved_page_recovery_uses_one_label_only(factory) -> None:
    with factory() as session, session.begin():
        one = _company(session, "Одна")
        many = _company(session, "Много")
        _card(
            session,
            one,
            "2gis",
            {"name": "Одна", "city": "Алматы"},
            {"category": {"reason": "unknown_label"}, "_collection": {"raw_path": ONE}},
            "https://2gis.kz/almaty/firm/9",
        )
        gis = list_platform_rows(session, "2gis")["items"][0]
        assert gis["category"] == "Не указана"
        yandex_one = _card(
            session,
            one,
            "yandex",
            {"name": "Одна"},
            {"category": {"reason": "unknown_label"}, "_collection": {"raw_path": ONE}},
            "https://yandex.kz/maps/org/9",
        )
        yandex_many = _card(
            session,
            many,
            "yandex",
            {"name": "Много"},
            {"category": {"reason": "unknown_label"}, "_collection": {"raw_path": MANY}},
            "https://yandex.kz/maps/org/10",
        )
        from zhol.agents.scout.platform_rows import _recover_single_label

        assert _recover_single_label(yandex_one) == "Салон красоты"
        assert _recover_single_label(yandex_many) is None
        assert yandex_one.fields.get("category") is None
        assert yandex_many.fields.get("category") is None
