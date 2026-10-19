from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.collect import Identity, collect_matched
from zhol.agents.scout.enrich import resolve_contacts
from zhol.agents.scout.maps import run_maps
from zhol.agents.scout.match import Candidate, candidates_from_document, choose
from zhol.agents.scout.persist import store_listed_services
from zhol.browser.fixtures import LoadedPage
from zhol.browser.maps.visible import candidates_from_links, parse_card, place_id
from zhol.browser.page import AccessRestricted, PageDocument
from zhol.cli import main
from zhol.db.models import CardSnapshot, Company, CompanyPlatformId, CompanyProfile, HumanTask
from zhol.export.mailing import COLUMNS, build_row, load_categories, write_xlsx

BADI = "70000001025141518"
VELARA = "70000001109698727"
CARD = """Badi
Салон красоты
4.8
84 оценки
24 отзыва
69 фото
Микрорайон Самал-2, 32
Алматы
Ежедневно с 09:00 до 22:00
+7 701 949 60 08
WhatsApp +7 701 949 60 08
Instagram @badisalon
badi@example.com
Описание
Салон в Самале
Услуги
Стрижка
5 000 ₸
Похожие организации
+7 700 000 00 00
"""
HREFS = [
    "https://wa.me/77019496008",
    "https://instagram.com/badisalon",
    "https://badi.example",
    "https://yandex.kz/maps/org/badi/166357199721/",
]


def _snap(**fields: object):
    data = dict(fields)
    return type(
        "Snap",
        (),
        {
            "fields": data,
            "url": "https://example.test/card",
            "screenshot_path": None,
            "status": "collected",
            "match_status": data.pop("match_status", "found"),
            "collected_at": None,
            "fail_reason": None,
            "field_sources": {},
        },
    )()


def test_public_text_keeps_explicit_contacts_and_drops_a_neighbour_phone() -> None:
    fields, notes = parse_card(CARD, name="Badi", hrefs=HREFS, platform="yandex")
    assert fields["name"] == "Badi"
    assert fields["category"] == "Салон красоты"
    assert fields["phones"] == "+7 701 949 60 08"
    assert fields["whatsapp"] == "+7 701 949 60 08"
    assert fields["instagram"] == "badisalon"
    assert fields["email"] == "badi@example.com"
    assert fields["website"] == "https://badi.example"
    assert fields["rating"] == "4.8"
    assert fields["reviews_count"] == "24"
    assert fields["photos_count"] == "69"
    assert fields["city"] == "Алматы"
    assert "Самал-2" in fields["address"]
    assert fields["hours"].startswith("Ежедневно")
    assert fields["description"] == "Салон в Самале"
    assert notes["service_names"] == "Стрижка"
    bare, bare_notes = parse_card("Badi\nСалон красоты\n+7 701 949 60 08\n", name="Badi", hrefs=[], platform="google")
    assert "whatsapp" not in bare and bare_notes["whatsapp"] == "not_shown"
    assert "instagram" not in bare and bare_notes["instagram"] == "not_shown"
    broken, broken_notes = parse_card(
        "Badi\n+7 (701) 949-60-08\n",
        name="Badi",
        hrefs=["https://www.instagram.com/`3w4r", "https://wa.me/77019496008"],
        platform="yandex",
    )
    assert "instagram" not in broken and broken_notes["instagram"] == "not_shown"
    assert broken["whatsapp"] == "+77019496008"
    assert "email" not in bare and bare_notes["email"] == "not_shown"
    assert "photos_count" not in bare and bare_notes["photos_count"] == "not_shown"
    assert "description" not in bare and bare_notes["description"] == "not_shown"


def test_google_review_count_comes_from_the_rating_line_only() -> None:
    fields, notes = parse_card("Velara\nСалон красоты\n4.6\n(18)\nАлматы\n", name="Velara", hrefs=[], platform="google")
    assert fields["rating"] == "4.6"
    assert fields["reviews_count"] == "18"
    assert "photos_count" not in fields and notes["photos_count"] == "not_shown"
    assert place_id("yandex", "https://yandex.kz/maps/org/badi/166357199721/") == "166357199721"
    assert place_id("google", "https://www.google.com/maps/place/Badi/data=!1s0x38836f:0xabc") == "0x38836f:0xabc"
    assert place_id("google", "https://maps.google.com/?place_id=ChIJbadi") == "ChIJbadi"


def test_floor_marker_and_social_links_do_not_hide_the_card() -> None:
    text = "\n".join(
        [
            "Velara beauty",
            "Салон красоты",
            "Ещё нет отзывов",
            "Алматы, просп. Назарбаева, 226",
            "этаж 1",
            "Адрес",
            "Алматы, просп. Назарбаева, 226этаж 1",
            "zapis.kz",
            "Как добраться",
            "Улица Жолдасбекова",
        ]
    )
    hrefs = [
        "https://zapis.kz/salon/1",
        "https://www.facebook.com/velara",
        "https://t.me/mapsyandex",
        "https://yastatic.net/font.woff2",
        "https://ads.adfox.ru/pixel",
        "https://instagram.com/velara.beauty.studio",
        "https://wa.me/77718086666",
    ]
    fields, notes = parse_card(text, name="Velara beauty", hrefs=hrefs, platform="yandex")
    assert fields["address"] == "Алматы, просп. Назарбаева, 226"
    assert notes["address"] == "shown"
    assert fields["reviews_count"] == "0"
    assert fields["website"] == "https://zapis.kz/salon/1"
    assert fields["instagram"] == "velara.beauty.studio"
    assert fields["whatsapp"] == "+77718086666"


def test_a_photo_count_is_not_the_business_name() -> None:
    links = [("https://yandex.kz/maps/org/badi/166357199721/", "69\nBadi\nмикрорайон Самал-2, 32\n+7 (701) 949-60-08")]
    candidates = candidates_from_links("yandex", links)
    assert len(candidates) == 1
    assert candidates[0]["fields"]["name"] == "Badi"
    identity = Identity("Badi", "Микрорайон Самал-2, 32", [])
    found = candidates_from_document(
        PageDocument(page="search", platform="yandex", external_id=None, fields={}, hrefs={}, candidates=candidates)
    )
    assert choose(identity, found).verdict == "same"


def test_duplicate_links_collapse_and_a_shared_phone_is_not_merged() -> None:
    links = [
        ("https://yandex.kz/maps/org/badi/1/", "Badi\nмикрорайон Самал-2, 32\n+7 701 949 60 08"),
        ("https://yandex.kz/maps/org/badi/1/", "Badi\nмикрорайон Самал-2, 32\n+7 701 949 60 08"),
        ("https://yandex.kz/maps/org/badi/2/", "Badi\nпроспект Назарбаева, 226\n+7 701 949 60 08"),
    ]
    candidates = candidates_from_links("yandex", links)
    assert len(candidates) == 2
    identity = Identity("Badi", "Микрорайон Самал-2, 32", ["+7 701 949 60 08"])
    same = Candidate("1", "1", links[0][0], "Badi", "микрорайон Самал-2, 32", ["+7 701 949 60 08"])
    other = Candidate("2", "2", links[2][0], "Badi", "проспект Назарбаева, 226", ["+7 701 949 60 08"])
    assert choose(identity, [same]).verdict == "same"
    assert choose(identity, [other]).verdict == "ambiguous"
    assert choose(identity, [same, other]).verdict == "same"
    assert choose(identity, [other, other]).verdict == "ambiguous"
    mall = Candidate("m", "9", "https://yandex.example/m", "Velara", "проспект Назарбаева, 226", [])
    venue = Identity("Velara", "МФК Samal Mall Проспект Назарбаева, 226", [])
    assert choose(venue, [mall]).verdict == "same"
    stranger = Candidate("s", "8", "https://yandex.example/s", "Самал Молл", "проспект Назарбаева, 226", [])
    assert choose(venue, [stranger]).verdict == "not_match"
    listed = Candidate(
        "v",
        "53907907781",
        "https://yandex.kz/maps/org/velara_beauty/53907907781/",
        "Velara beauty",
        "Алматы, просп. Назарбаева, 226этаж 1",
        ["+7 (771) 808-66-66"],
    )
    assert choose(venue, [listed]).verdict == "same"
    assert choose(venue, [listed]).phone_match is False


def test_contact_priority_fills_a_blank_and_reports_a_conflict() -> None:
    snapshots = {
        "2gis": _snap(phones=["+7 701 949 60 08"], whatsapp=None, link="https://zapis.kz/salon/437", email="badi@example.com"),
        "yandex": _snap(phones=["8 701 949 60 08"], whatsapp="+7 701 949 60 08", website="https://badi.example"),
        "google": _snap(phones=["+7 702 000 00 00"], website="https://other.example", email=None, match_status="found"),
    }
    chosen, conflicts = resolve_contacts(snapshots, company_whatsapp="+7 777 000 00 00", company_instagram="@badisalon")
    assert chosen["phone"] == "+7 701 949 60 08"
    assert any(item.startswith("contact phone:") and "kept 2gis" in item for item in conflicts)
    assert chosen["whatsapp"] == "+7 701 949 60 08"
    assert chosen["website"] == "https://other.example"
    assert any(item.startswith("contact website:") for item in conflicts)
    assert chosen["email"] == "badi@example.com"
    assert chosen["instagram"] == "@badisalon"
    blocked = {
        "2gis": _snap(phones=["+7 701 949 60 08"]),
        "yandex": _snap(phones=["+7 702 000 00 00"], match_status="other_address"),
    }
    quiet, quiet_conflicts = resolve_contacts(blocked)
    assert quiet["phone"] == "+7 701 949 60 08"
    assert quiet_conflicts == []
    assert resolve_contacts({"2gis": _snap(phones=None)}, company_whatsapp="+7 701 949 60 08")[0]["whatsapp"] == "+7 701 949 60 08"


def test_excel_keeps_forty_columns_and_shows_the_conflict(tmp_path: Path) -> None:
    table = load_categories()
    row = build_row(
        company_id=BADI,
        name="Badi",
        city="Алматы",
        niche="салон",
        source="test",
        whatsapp=None,
        instagram=None,
        snapshots={
            "2gis": _snap(phones=["+7 701 949 60 08"], category="Салон красоты", photos_count=146, description="в 2ГИС"),
            "google": _snap(phones=["+7 702 000 00 00"], photos_count=4, rating=4.6, reviews_count=18, description="в Google"),
            "yandex": _snap(rating=4.9, reviews_count=24, photos_count=69, hours="Ежедневно с 09:00 до 22:00"),
        },
        services=[{"name": "Стрижка", "needs_confirmation": False}],
        score=70,
        discrepancies=[],
        categories=table,
    )
    path = tmp_path / "mailing.xlsx"
    assert write_xlsx([row], path) == 1
    sheet = load_workbook(path).active
    assert tuple(cell.value for cell in sheet[1]) == COLUMNS
    assert len(COLUMNS) == 40
    header = {cell.value: cell.column for cell in sheet[1]}
    assert row["Phone"] == "+7 701 949 60 08"
    assert sheet.cell(2, header["Phone"]).value == "'+7 701 949 60 08"
    assert sheet.cell(2, header["Google Photo Count"]).value == 4
    assert sheet.cell(2, header["Google Description"]).value == "в Google"
    assert sheet.cell(2, header["Yandex Opening Hours"]).value == "Ежедневно с 09:00 до 22:00"
    assert "contact phone:" in sheet.cell(2, header["Identified Problems"]).value
    assert row["Services / Activities"] == "Стрижка"


def test_uncertain_match_is_escalated_and_a_rerun_does_not_insert_again(migrated_engine: Engine) -> None:
    factory = sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)
    identity = Identity("Badi", "Микрорайон Самал-2, 32", ["+7 701 949 60 08"])
    search = PageDocument(
        page="search",
        platform="yandex",
        external_id=None,
        fields={},
        hrefs={},
        candidates=[
            {
                "candidate_id": "1",
                "external_id": "166357199721",
                "url": "https://yandex.kz/maps/org/badi/166357199721/",
                "fields": {"name": "Badi", "address": "микрорайон Самал-2, 32", "phones": "+7 701 949 60 08"},
            }
        ],
    )
    fields, notes = parse_card(CARD, name="Badi", hrefs=HREFS, platform="yandex")
    card = PageDocument(
        page="card",
        platform="yandex",
        external_id="166357199721",
        fields=fields,
        hrefs={},
        notes=notes,
    )
    search_page = LoadedPage("https://yandex.kz/maps/?text=badi", "", "search.html", "search.html", "search.png", "search", search)
    card_page = LoadedPage(search.candidates[0]["url"], "", "card.html", "card.html", "card.png", "card", card)

    class Ready:
        def search(self, platform: str, id_2gis: str, url: str) -> LoadedPage:
            return search_page

        def card(self, platform: str, external_id: str, url: str) -> LoadedPage:
            return card_page

    conflict = PageDocument(
        page="search",
        platform="google",
        external_id=None,
        fields={},
        hrefs={},
        candidates=[
            {
                "candidate_id": "a",
                "external_id": "ChIJone",
                "url": "https://www.google.com/maps/place/a",
                "fields": {"name": "Badi", "address": "проспект Назарбаева, 226", "phones": "+7 701 949 60 08"},
            }
        ],
    )
    conflict_page = LoadedPage("https://www.google.com/maps/search/?query=badi", "", "g.html", "g.html", "g.png", "search", conflict)

    class Conflict:
        def search(self, platform: str, id_2gis: str, url: str) -> LoadedPage:
            return conflict_page

        def card(self, platform: str, external_id: str, url: str) -> LoadedPage:
            raise AssertionError("ambiguous card must stay closed")

    from zhol.browser.config import load_platforms

    cfg = load_platforms()
    with factory() as session, session.begin():
        company = Company(id_2gis=BADI, name="Badi", source="test", city="Алматы", whatsapp="+7 701 949 60 08")
        session.add(company)
        session.flush()
        company_id = company.id
        session.add(
            CardSnapshot(
                company_id=company_id,
                platform="2gis",
                match_status="found",
                url="https://2gis.kz/almaty/firm/" + BADI,
                fields={"name": "Badi", "rating": 4.8},
                field_sources={},
                purpose="prospect",
                collected_at=__import__("datetime").datetime(2026, 10, 11, tzinfo=__import__("datetime").timezone.utc),
                collected_by="scout",
                status="collected",
            )
        )
        for _ in range(2):
            done = collect_matched(
                session,
                company=company,
                identity=identity,
                purpose="prospect",
                collection_key=f"e3-maps:{BADI}",
                pages=Ready(),  # type: ignore[arg-type]
                cfg=cfg,
                platforms=("yandex",),
            )
            assert done[0].verdict in ("same", "reused")
        missed = collect_matched(
            session,
            company=company,
            identity=identity,
            purpose="prospect",
            collection_key=f"e3-maps:{BADI}",
            pages=Conflict(),  # type: ignore[arg-type]
            cfg=cfg,
            platforms=("google",),
        )
        assert missed[0].verdict == "ambiguous" and missed[0].snapshot is None and missed[0].hold == "ambiguous_match"
        assert store_listed_services(session, company_id, ["Стрижка"], source_ref="field:yandex.services", forbidden=["Badi"])
        assert "не заменён" in store_listed_services(
            session, company_id, ["Маникюр"], source_ref="field:google.services", forbidden=["Badi"]
        )
    with factory() as session:
        yandex = session.scalars(select(CardSnapshot).where(CardSnapshot.company_id == company_id, CardSnapshot.platform == "yandex")).all()
        google = session.scalars(select(CardSnapshot).where(CardSnapshot.company_id == company_id, CardSnapshot.platform == "google")).all()
        twogis = session.scalars(select(CardSnapshot).where(CardSnapshot.company_id == company_id, CardSnapshot.platform == "2gis")).all()
        assert len(yandex) == 1 and yandex[0].fields["email"] == "badi@example.com"
        assert yandex[0].fields["whatsapp"] == "+7 701 949 60 08"
        assert session.get(CompanyPlatformId, (company_id, "yandex")).external_id == "166357199721"
        assert google == []
        assert len(twogis) == 1 and twogis[0].fields["rating"] == 4.8
        assert session.scalar(select(func.count()).select_from(Company)) == 1
        assert session.scalar(select(HumanTask).where(HumanTask.reason == "ambiguous_match")) is not None
        assert session.scalar(select(HumanTask).where(HumanTask.reason == "service_conflict")) is not None
        profile = session.scalar(select(CompanyProfile).where(CompanyProfile.company_id == company_id))
        assert [item["name"] for item in profile.data["services"]] == ["Стрижка"]


def test_robot_check_stops_the_next_page(tmp_path: Path) -> None:
    opened: list[str] = []

    class Opener:
        def search(self, platform: str, id_2gis: str, url: str, folder: Path) -> LoadedPage:
            opened.append(platform + ":" + id_2gis)
            raise AccessRestricted("access=robot_check", screenshot=str(folder / "screenshot.png"), url=url, access="robot_check", outcome="access_blocked")

        def card(self, platform: str, external_id: str, url: str, folder: Path) -> LoadedPage:
            raise AssertionError(platform)

    code = run_maps(
        [BADI, VELARA],
        output=tmp_path / "out",
        report_path=tmp_path / "e3.md",
        identities={
            BADI: Identity("Badi", "Микрорайон Самал-2, 32", ["+7 701 949 60 08"]),
            VELARA: Identity("Velara", "Проспект Назарбаева, 226", []),
        },
        opener=Opener(),
        exporter=lambda: "not-written",
    )
    text = (tmp_path / "e3.md").read_text(encoding="utf-8")
    assert code == 2
    assert opened == ["yandex:" + BADI]
    assert "Не открывалась" in text
    assert "файл не перезаписан" in text


def test_maps_command_rejects_the_wrong_batch(capsys, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.delenv("ZHOL_SCOUT_LIVE", raising=False)
    two = tmp_path / "two.txt"
    two.write_text(BADI + "\n" + VELARA + "\n", encoding="utf-8")
    assert main(["scout", "--maps", "--ids-file", str(two)]) == 2
    assert "requires --live" in capsys.readouterr().err
    three = tmp_path / "three.txt"
    three.write_text("1\n2\n3\n", encoding="utf-8")
    monkeypatch.setenv("ZHOL_SCOUT_LIVE", "1")
    assert main(["scout", "--live", "--maps", "--ids-file", str(three)]) == 2
    assert "exactly 2" in capsys.readouterr().err
