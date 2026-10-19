"""E4.1 saved-page discovery. No live catalog requests."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.discover import run_discovery
from zhol.browser.discovery.engine import walk_saved
from zhol.browser.discovery.normalize import (
    normalize_category,
    verified_email,
    verified_instagram,
    verified_phone,
    verified_website,
    verified_whatsapp,
)
from zhol.browser.discovery.session import PlaywrightSavedSession
from zhol.browser.discovery.types import DiscoveryRefused, RunControl
from zhol.db.models import CardSnapshot, Company
from zhol.etalon import SOURCE, import_etalon
from zhol.export.mailing import COLUMNS
from zhol.guards import saved_discovery_only

ROOT = Path("tests/fixtures/discovery/almaty-salon")
PAGES = ROOT / "pages"
ETALON = Path("source/etalon-50.csv")
SALON = "Салон красоты"
FIRMS = {
    "70000001000000901",
    "70000001000000902",
    "70000001000000903",
    "70000001000000904",
}


def _factory(engine: Engine):
    return sessionmaker(bind=engine, expire_on_commit=False, future=True)


def _sheet(path: Path):
    book = load_workbook(path)
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    rows = [dict(zip(header, cells)) for cells in sheet.iter_rows(min_row=2, values_only=True)]
    return header, rows


def _walk(tmp_path: Path, **kwargs):
    options = dict(
        city="Алматы",
        category=SALON,
        platforms=("2gis", "yandex", "google"),
        source="fixture",
        evidence_dir=tmp_path / "evidence",
    )
    options.update(kwargs)
    return walk_saved(ROOT, **options)


def test_saved_search_finds_firms_pages_and_scroll(tmp_path: Path) -> None:
    result = _walk(tmp_path)
    firms = {branch.firm_id for branch in result.branches}
    assert firms == FIRMS
    assert "70000001000000999" not in firms
    assert "70000001000000906" not in firms
    assert "70000001000000905" not in firms
    assert result.progress["pages"] >= 2
    assert result.progress["scrolls"] == 1
    assert result.progress["duplicates"] >= 1
    assert result.progress["blocked"][0]["access"] == "robot_check"
    assert result.progress["missing"]
    assert any(item["reason"] == "no_firm_url" for item in result.progress["skipped"])
    north = next(branch for branch in result.branches if branch.firm_id == "70000001000000901")
    south = next(branch for branch in result.branches if branch.firm_id == "70000001000000903")
    west = next(branch for branch in result.branches if branch.firm_id == "70000001000000904")
    assert north.name == "North Beauty"
    assert north.category == SALON
    assert north.phones == ("77010000901",)
    assert north.whatsapp == "77010000901"
    assert north.instagram == "north.beauty"
    assert north.website == "https://north.example/salon"
    assert north.email == "north@example.com"
    assert north.platforms["2gis"].seen_at.startswith("2026-10-11")
    assert Path(north.platforms["2gis"].screenshot_path).is_file()
    assert set(north.platforms) == {"2gis", "yandex", "google"}
    assert south.subcategory == "Парикмахерская"
    assert west.phones == ()
    second = next(branch for branch in result.branches if branch.firm_id == "70000001000000902")
    assert second.phones == ("77010000902",)
    assert second.phone_conflict is True
    assert "yandex" not in second.platforms
    assert any(item["platform"] == "yandex" and item["reason"] == "ambiguous" for item in second.alternatives)
    assert second.address != north.address


def test_pause_and_resume_walk_does_not_repeat_a_page(tmp_path: Path) -> None:
    paused = _walk(tmp_path, platforms=("2gis",), control=RunControl(pause_after_pages=1))
    assert paused.status == "paused"
    assert {branch.firm_id for branch in paused.branches} == {"70000001000000901", "70000001000000902"}
    resumed = walk_saved(
        ROOT,
        city="Алматы",
        category=SALON,
        platforms=("2gis",),
        source="fixture",
        evidence_dir=tmp_path / "more",
        resume_state=paused.progress,
        control=RunControl(),
    )
    assert {branch.firm_id for branch in resumed.branches} == FIRMS
    assert "70000001000000999" not in {branch.firm_id for branch in resumed.branches}


def test_live_catalog_is_refused_before_a_page_opens(tmp_path: Path) -> None:
    class Page:
        def set_content(self, html: str, wait_until: str | None = None) -> None:
            raise AssertionError("a page was loaded")

        def goto(self, url: str, **kwargs) -> None:
            raise AssertionError(url)

    with pytest.raises(PermissionError):
        walk_saved(
            ROOT,
            city="Алматы",
            category=SALON,
            platforms=("2gis",),
            source="live",
            evidence_dir=tmp_path,
            page=Page(),
        )
    with pytest.raises(PermissionError):
        saved_discovery_only("live", "yandex")
    saved_discovery_only("fixture", "2gis")
    saved_discovery_only("manual", "google")


def test_playwright_session_refuses_a_website() -> None:
    class Page:
        def set_content(self, html: str, wait_until: str | None = None) -> None:
            raise AssertionError(html)

        def screenshot(self, path: str, full_page: bool = True) -> None:
            raise AssertionError(path)

        def inner_text(self, selector: str) -> str:
            raise AssertionError(selector)

    with pytest.raises(DiscoveryRefused):
        PlaywrightSavedSession(Page()).show("<html></html>", "https://2gis.kz/almaty/search", Path("x.png"))


def test_category_and_contacts_stay_empty_when_unverified() -> None:
    sector, category, subcategory = normalize_category("Beauty salon")
    assert (sector, category, subcategory) == ("Красота", SALON, None)
    assert normalize_category("салон") == (None, None, None)
    assert normalize_category("Парикмахерская")[1:] == (SALON, "Парикмахерская")
    assert normalize_category("") == (None, None, None)
    assert verified_phone("123") is None
    assert verified_phone("8 701 000 09 01") == "77010000901"
    assert verified_whatsapp("https://example.com", "call me") is None
    assert verified_instagram(None, "not a handle") is None
    assert verified_website("https://yandex.kz/maps/org/1", None) is None
    assert verified_email("north") is None


def test_pause_leaves_a_header_and_resume_fills_excel(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "paused.xlsx"
    paused = run_discovery(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis", "yandex", "google"),
        limit=None,
        out=out,
        source_mode="fixture",
        root=ROOT,
        fixtures_dir=PAGES,
        control=RunControl(pause_after_pages=1),
    )
    assert paused.completed == 0
    assert paused.pending == 2
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert rows == []
    resumed = run_discovery(
        factory,
        city=None,
        category=None,
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="fixture",
        root=ROOT,
        resume=paused.batch_id,
        control=RunControl(),
    )
    assert resumed.completed == 4
    assert resumed.pending == 0
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert {row["Company ID"] for row in rows} == FIRMS
    assert not (out.parent / (out.name + ".partial")).exists()


def test_partial_excel_and_resume_do_not_duplicate_snapshots(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "partial.xlsx"
    first = run_discovery(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="fixture",
        root=ROOT,
        fixtures_dir=PAGES,
        stop_after=1,
    )
    assert first.completed == 1
    assert first.pending == 3
    _header, rows = _sheet(out)
    assert len(rows) == 1
    assert rows[0]["Company ID"] == "70000001000000901"
    assert rows[0]["Business Category"] == SALON
    assert rows[0]["2GIS Review Count"] == 0
    assert rows[0]["Source"] == "saved_pages"
    with factory() as session:
        before = session.scalars(select(CardSnapshot.id).order_by(CardSnapshot.id)).all()
    second = run_discovery(
        factory,
        city=None,
        category=None,
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="fixture",
        root=ROOT,
        resume=first.batch_id,
    )
    assert second.completed == 4
    assert second.duplicate == 1
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert len(rows) == 4
    south = next(row for row in rows if row["Company ID"] == "70000001000000903")
    assert south["Business Category"] == SALON
    assert south["Business Subcategory"] == "Парикмахерская"
    with factory() as session:
        after = session.scalars(select(CardSnapshot.id).order_by(CardSnapshot.id)).all()
        assert session.scalar(select(func.count()).select_from(CardSnapshot).where(CardSnapshot.company_id.is_not(None))) == 4
    assert before[0] in after
    assert len(after) == len(before) + 3


def test_reference_snapshots_stay_in_place(tmp_path: Path, migrated_engine: Engine) -> None:
    if not ETALON.is_file():
        pytest.skip("Файл source/etalon-50.csv не входит в публичный репозиторий.")
    factory = _factory(migrated_engine)
    with factory() as session, session.begin():
        import_etalon(session, ETALON)
    with factory() as session:
        before = set(session.scalars(select(CardSnapshot.id).join(Company).where(Company.source == SOURCE)).all())
        companies = session.scalar(select(func.count()).select_from(Company).where(Company.source == SOURCE))
    run_discovery(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis",),
        limit=None,
        out=tmp_path / "reference.xlsx",
        source_mode="fixture",
        root=ROOT,
        fixtures_dir=PAGES,
    )
    with factory() as session:
        after = set(session.scalars(select(CardSnapshot.id).join(Company).where(Company.source == SOURCE)).all())
        assert session.scalar(select(func.count()).select_from(Company).where(Company.source == SOURCE)) == companies
        assert session.scalar(select(Company.id).where(Company.id_2gis == "70000001000000999")) is None
    assert after == before
    assert companies == 50


def test_manual_pages_use_the_same_excel(tmp_path: Path, migrated_engine: Engine) -> None:
    root = tmp_path / "manual"
    (root / "2gis").mkdir(parents=True)
    (root / "pages" / "2gis").mkdir(parents=True)
    (root / "manifest.yaml").write_text(
        "city: Алматы\ncategory: Салон красоты\nstarts:\n  2gis: 2gis/page-1.html\n",
        encoding="utf-8",
    )
    (root / "2gis" / "page-1.html").write_text(
        "<html><body><section data-page=\"search\" data-platform=\"2gis\">"
        "<a data-candidate=\"1\" href=\"https://2gis.kz/almaty/firm/70000001000000901\">"
        "<span data-field=\"name\">North Beauty</span>"
        "<span data-field=\"address\">улица Абая, 10</span>"
        "<span data-field=\"category\">Салон красоты</span>"
        "</a></section></body></html>",
        encoding="utf-8",
    )
    (root / "pages" / "blank.png").write_bytes((PAGES / "blank.png").read_bytes())
    (root / "pages" / "2gis" / "70000001000000901.html").write_text(
        (PAGES / "2gis" / "70000001000000901.html").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    factory = _factory(migrated_engine)
    out = tmp_path / "manual.xlsx"
    counts = run_discovery(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="manual",
        root=root,
        fixtures_dir=root / "pages",
    )
    assert counts.completed == 1
    _header, rows = _sheet(out)
    assert rows[0]["Source"] == "manual_browser"
    assert rows[0]["Company Name"] == "North Beauty"
    with factory() as session:
        company = session.scalar(select(Company).where(Company.id_2gis == "70000001000000901"))
        assert company is not None
        assert company.source == "manual_browser"
