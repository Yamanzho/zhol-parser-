from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.batch import run_batch
from zhol.agents.scout.catalog import BatchError, load_catalog, select_companies
from zhol.browser.fixtures import FixtureStore
from zhol.browser.page import AccessRestricted
from zhol.cli import main
from zhol.db.models import CardSnapshot, CollectionBatchItem, Company
from zhol.etalon import SOURCE, import_etalon
from zhol.export.mailing import COLUMNS, write_xlsx

CATALOG = Path("tests/fixtures/catalog.yaml")
ETALON = Path("source/etalon-50.csv")
SALON = "Салон красоты"


def _factory(engine: Engine):
    return sessionmaker(bind=engine, expire_on_commit=False, future=True)


def _sheet(path: Path):
    book = load_workbook(path)
    sheet = book.active
    header = [cell.value for cell in sheet[1]]
    rows = []
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, cells)))
    return header, rows


class _BlockGoogle(FixtureStore):
    def search(self, platform: str, id_2gis: str, url: str):
        if platform == "google":
            raise AccessRestricted("robot check", access="robot_check", outcome="stopped")
        return super().search(platform, id_2gis, url)


def test_category_keywords_and_two_cities() -> None:
    rows = load_catalog(CATALOG)
    almaty = select_companies(rows, city="алматы", category=SALON)
    astana = select_companies(rows, city="Астана", category="Ресторан")
    assert len(almaty) == 10
    assert len(astana) == 2
    assert {row.city for row in almaty} == {"Алматы"}
    assert {row.category for row in astana} == {"Ресторан"}
    try:
        select_companies(rows, city="Алматы", category="салон")
    except BatchError as exc:
        assert "ambiguous" in str(exc)
    else:
        raise AssertionError("ambiguous category was accepted")


def test_limit_is_optional_and_does_not_invent_companies(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "limited.xlsx"
    limited = run_batch(
        factory,
        city="Астана",
        category="Ресторан",
        platforms=("2gis",),
        limit=1,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
    )
    assert limited.completed == 1
    assert limited.pending == 0
    assert limited.rows == 1
    unlimited = run_batch(
        factory,
        city="Астана",
        category="Ресторан",
        platforms=("2gis",),
        limit=None,
        out=tmp_path / "all.xlsx",
        source_mode="fixture",
        catalog=CATALOG,
    )
    assert unlimited.completed == 2
    capped = run_batch(
        factory,
        city="Астана",
        category="Ресторан",
        platforms=("2gis",),
        limit=100,
        out=tmp_path / "capped.xlsx",
        source_mode="fixture",
        catalog=CATALOG,
    )
    assert capped.completed == 2


def test_resume_does_not_duplicate_snapshots(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "resume.xlsx"
    first = run_batch(
        factory,
        city="Астана",
        category="Ресторан",
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
        stop_after=1,
    )
    assert first.completed == 1
    assert first.pending == 1
    assert first.skipped == 1
    with factory() as session:
        before = session.scalars(select(CardSnapshot.id).order_by(CardSnapshot.id)).all()
        stored = session.get(CardSnapshot, before[0])
        fields = dict(stored.fields)
    second = run_batch(
        factory,
        city=None,
        category=None,
        platforms=("2gis",),
        limit=None,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
        resume=first.batch_id,
    )
    assert second.completed == 2
    assert second.duplicate == 1
    assert second.pending == 0
    with factory() as session:
        after = session.scalars(select(CardSnapshot.id).order_by(CardSnapshot.id)).all()
        kept = session.get(CardSnapshot, before[0])
        assert dict(kept.fields) == fields
        assert session.scalar(select(func.count()).select_from(Company).where(Company.id_2gis == "70000001000000041")) == 1
    assert before[0] in after
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert len(rows) == 2
    assert not (out.parent / (out.name + ".partial")).exists()


def test_blocked_platform_does_not_become_a_listing(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "blocked.xlsx"
    counts = run_batch(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis", "yandex", "google"),
        limit=1,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
        pages=_BlockGoogle(),
    )
    assert counts.completed == 1
    _header, rows = _sheet(out)
    row = rows[0]
    assert row["Google URL"] is None
    assert row["Google Rating"] is None
    assert "google: blocked" in row["Identified Problems"]
    assert row["2GIS URL"]
    assert row["Source"] == "fixture_catalog"


def test_ambiguous_match_is_not_merged(tmp_path: Path, migrated_engine: Engine) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "ambiguous.xlsx"
    counts = run_batch(
        factory,
        city="Алматы",
        category="Парикмахерская",
        platforms=("2gis", "yandex"),
        limit=None,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
    )
    assert counts.completed == 1
    _header, rows = _sheet(out)
    row = rows[0]
    assert "yandex: ambiguous" in row["Identified Problems"]
    assert "google: not_checked" in row["Identified Problems"]
    assert row["Yandex URL"] is None
    assert row["Phone"] == "'+7 701 000 00 18"
    assert row["Google Review Count"] is None


def test_missing_page_stays_empty(tmp_path: Path, migrated_engine: Engine) -> None:
    catalog = tmp_path / "catalog.yaml"
    catalog.write_text(
        "companies:\n"
        "  - {id_2gis: '70000001000000099', name: 'Missing', city: 'Алматы', category: 'Салон красоты'}\n",
        encoding="utf-8",
    )
    factory = _factory(migrated_engine)
    out = tmp_path / "missing.xlsx"
    counts = run_batch(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis", "yandex", "google"),
        limit=None,
        out=out,
        source_mode="fixture",
        catalog=catalog,
    )
    assert counts.failed == 1
    assert counts.completed == 0
    _header, rows = _sheet(out)
    row = rows[0]
    assert row["Phone"] is None
    assert row["2GIS URL"] is None
    assert row["2GIS Review Count"] is None
    assert "2gis: not_found" in row["Identified Problems"]
    assert "yandex: not_checked" in row["Identified Problems"]
    assert "google: not_checked" in row["Identified Problems"]
    assert row["Verification Status"] == "not_found"


def test_fixture_batch_of_ten_keeps_the_reference_set(tmp_path: Path, migrated_engine: Engine) -> None:
    if not ETALON.is_file():
        pytest.skip("Файл source/etalon-50.csv не входит в публичный репозиторий.")
    factory = _factory(migrated_engine)
    with factory() as session, session.begin():
        import_etalon(session, ETALON)
    with factory() as session:
        before = set(session.scalars(select(CardSnapshot.id).join(Company).where(Company.source == SOURCE)).all())
        companies = session.scalar(select(func.count()).select_from(Company).where(Company.source == SOURCE))
    out = tmp_path / "ten.xlsx"
    counts = run_batch(
        factory,
        city="Алматы",
        category=SALON,
        platforms=("2gis", "yandex", "google"),
        limit=10,
        out=out,
        source_mode="fixture",
        catalog=CATALOG,
    )
    assert counts.completed == 10
    assert counts.failed == 0
    assert counts.duplicate == 0
    assert counts.pending == 0
    assert counts.rows == 10
    with factory() as session:
        after = set(session.scalars(select(CardSnapshot.id).join(Company).where(Company.source == SOURCE)).all())
        assert session.scalar(select(func.count()).select_from(Company).where(Company.source == SOURCE)) == companies
    assert after == before
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert len(rows) == 10
    saved = {row["Company ID"]: row for row in rows}
    assert set(saved) == {row.id_2gis for row in select_companies(load_catalog(CATALOG), city="Алматы", category=SALON)}
    one = saved["70000001000000031"]
    assert one["2GIS Review Count"] == 0
    assert one["Company Name"] == "Salon One"
    assert one["Business Category"] == SALON
    assert one["Phone"] == "'+7 701 000 00 31"
    assert "yandex: not_found" in one["Identified Problems"]
    assert "google: not_found" in one["Identified Problems"]
    assert one["Source"] == "fixture_catalog"
    assert one["Verification Status"] == "collected"
    again = write_xlsx(
        [
            {
                **{column: one.get(column) for column in COLUMNS},
                "Phone": "+7 701 000 00 31",
            }
        ],
        tmp_path / "again.xlsx",
    )
    assert again == 1
    _header, repeated = _sheet(tmp_path / "again.xlsx")
    assert repeated[0]["Phone"] == "'+7 701 000 00 31"
    assert repeated[0]["2GIS Review Count"] == 0


def test_excel_replace_keeps_the_previous_file_when_save_fails(tmp_path: Path, monkeypatch) -> None:
    dest = tmp_path / "mailing.xlsx"
    write_xlsx([{"Company ID": "1", "Company Name": "Kept"}], dest)
    original = dest.read_bytes()

    def fail_save(self, filename):
        raise OSError("disk full")

    monkeypatch.setattr(Workbook, "save", fail_save)
    try:
        write_xlsx([{"Company ID": "2", "Company Name": "New"}], dest)
    except OSError:
        pass
    else:
        raise AssertionError("save failure was ignored")
    assert dest.read_bytes() == original


def test_live_discovery_is_refused(capsys, monkeypatch) -> None:
    monkeypatch.delenv("ZHOL_SCOUT_LIVE", raising=False)
    code = main(
        ["scout-batch", "--city", "Алматы", "--category", SALON, "--source-mode", "live"],
        factory=object(),
    )
    assert code == 2
    monkeypatch.setenv("ZHOL_SCOUT_LIVE", "1")
    code = main(
        ["scout-batch", "--city", "Алматы", "--category", SALON, "--source-mode", "live"],
        factory=object(),
    )
    assert code == 2
    error = capsys.readouterr().err
    assert "unsupported" in error
    assert "2gis" in error and "yandex" in error and "google" in error


def test_cli_exports_the_astana_batch(tmp_path: Path, migrated_engine: Engine, capsys) -> None:
    factory = _factory(migrated_engine)
    out = tmp_path / "cli.xlsx"
    code = main(
        [
            "scout-batch",
            "--city",
            "Астана",
            "--category",
            "Ресторан",
            "--platforms",
            "2gis",
            "--out",
            str(out),
            "--source-mode",
            "fixture",
        ],
        factory=factory,
    )
    assert code == 0
    text = capsys.readouterr().out
    assert "collected 2, failed 0, blocked 0, skipped 0, duplicate 0, pending 0" in text
    header, rows = _sheet(out)
    assert header == list(COLUMNS)
    assert {row["Company Name"] for row in rows} == {"Astana Table", "Astana Grill"}
    assert all("yandex: not_checked" in row["Identified Problems"] for row in rows)
    with factory() as session:
        items = session.scalars(select(CollectionBatchItem)).all()
    assert len(items) == 2
    assert {item.status for item in items} == {"completed"}
