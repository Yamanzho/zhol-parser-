"""Bek audit import. The workbook under test is the 50-row manual audit."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session, sessionmaker

from zhol.audit.bek import company_plan, import_bek_audit
from zhol.audit.workbook import read_workbook
from zhol.cli import main
from zhol.db.models import (
    Audit,
    CardSnapshot,
    ChangeItem,
    ChangeLog,
    Company,
    CompanyPlatformId,
    CompanyProfile,
    HumanTask,
    ListingImprovementPlan,
)
from zhol.etalon import import_etalon
from zhol.timeutil import ALMATY

ROOT = Path(__file__).resolve().parents[1]
ETALON = ROOT / "source" / "etalon-50.csv"
if not ETALON.is_file():
    pytest.skip("Файл source/etalon-50.csv не входит в публичный репозиторий.", allow_module_level=True)
FINISHED = frozenset({"completed", "prepared", "applied", "done", "submitted"})


@pytest.fixture
def factory(migrated_engine: Engine) -> sessionmaker[Session]:
    return sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)


def _workbook(path: Path) -> Path:
    rows = list(csv.reader(ETALON.open(encoding="utf-8-sig")))
    book = Workbook()
    cover = book.active
    cover.title = "сводка"
    cover["A1"] = "Салоны Алматы, октябрь 2026"
    cover["A3"] = "Этот лист не таблица карточек"
    notes = book.create_sheet("заметки")
    notes.append(["комментарий"])
    notes.append(["ручной просмотр 30.09.2026"])
    sheet = book.create_sheet("карточки")
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _csv_rows() -> list[dict[str, str]]:
    return list(csv.DictReader(ETALON.open(encoding="utf-8-sig")))


def _company(session: Session, id_2gis: str) -> Company:
    company = session.scalar(select(Company).where(Company.id_2gis == id_2gis))
    assert company is not None
    return company


def _tasks(plan: dict, task_type: str, platform: str | None = None) -> list[dict]:
    found = [task for task in [*plan["automatic_preparation"], *plan["manual_work"]] if task["type"] == task_type]
    if platform is not None:
        found = [task for task in found if task["platform"] == platform]
    return found


def _assert_not_finished(plan: dict) -> None:
    assert plan["listing_modified"] is False
    for task in [*plan["automatic_preparation"], *plan["manual_work"]]:
        assert task["completed"] is False
        assert task["status"] not in FINISHED


def test_workbook_uses_the_detailed_sheet(tmp_path: Path) -> None:
    book = read_workbook(_workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx"))
    assert book.sheets == ("сводка", "заметки", "карточки")
    assert book.sheet == "карточки"
    assert len(book.rows) == 50
    assert book.rows[0].text("id_2gis") == "70000001066660465"


def test_import_preserves_scores_and_is_idempotent(factory, tmp_path: Path) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    expected = {row["id_2gis"]: row for row in _csv_rows()}
    with factory() as session, session.begin():
        import_etalon(session, ETALON)
        before = {
            row.id: dict(row.fields)
            for row in session.scalars(select(CardSnapshot).where(CardSnapshot.collected_by == "manual_import"))
        }
        first = import_bek_audit(session, path)
    assert first.sheet == "карточки"
    assert first.companies == 0
    assert first.snapshots == 150
    assert first.audits == 50
    assert first.plans == 50
    assert first.unmatched == []
    with factory() as session:
        audits = session.scalars(select(Audit).where(Audit.prompt_version == "bek-audit-import-1")).all()
        assert len(audits) == 50
        for audit in audits:
            company = session.get(Company, audit.company_id)
            row = expected[company.id_2gis]
            assert audit.score_total == int(row["балл_итого"])
            assert audit.score_groups == {
                "reviews": int(row["балл_отзывы"]),
                "prices": int(row["балл_цены"]),
                "photos": int(row["балл_фото"]),
                "contacts": int(row["балл_контакты"]),
            }
            assert company.source == "etalon_2026_09_30"
        for snapshot_id, fields in before.items():
            assert session.get(CardSnapshot, snapshot_id).fields == fields
        counts = (
            session.scalar(select(func.count()).select_from(Company)),
            session.scalar(select(func.count()).select_from(CardSnapshot)),
            session.scalar(select(func.count()).select_from(Audit)),
            session.scalar(select(func.count()).select_from(ListingImprovementPlan)),
        )
    with factory() as session, session.begin():
        second = import_bek_audit(session, path)
    assert (second.companies, second.snapshots, second.audits, second.plans) == (0, 0, 0, 0)
    with factory() as session:
        again = (
            session.scalar(select(func.count()).select_from(Company)),
            session.scalar(select(func.count()).select_from(CardSnapshot)),
            session.scalar(select(func.count()).select_from(Audit)),
            session.scalar(select(func.count()).select_from(ListingImprovementPlan)),
        )
    assert again == counts


def test_missing_google_unclaimed_photo_price_replies_and_nulls(factory, tmp_path: Path) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    with factory() as session, session.begin():
        import_bek_audit(session, path)
        montebello = _company(session, "70000001066660465")
        badi = _company(session, "70000001025141518")
        cozy = _company(session, "70000001054335834")
        style = _company(session, "70000001054653678")
        montebello_plan = session.scalar(
            select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == montebello.id)
        ).plan
        badi_plan = session.scalar(select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == badi.id)).plan
        cozy_plan = session.scalar(select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == cozy.id)).plan
        style_plan = session.scalar(select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == style.id)).plan
        google = session.scalar(
            select(CardSnapshot).where(
                CardSnapshot.company_id == montebello.id,
                CardSnapshot.platform == "google",
                CardSnapshot.collected_by == "bek_audit",
            )
        )
        twogis = session.scalar(
            select(CardSnapshot).where(
                CardSnapshot.company_id == montebello.id,
                CardSnapshot.platform == "2gis",
                CardSnapshot.collected_by == "bek_audit",
            )
        )
        blanks = [
            row
            for row in session.scalars(
                select(CardSnapshot).where(
                    CardSnapshot.platform == "google",
                    CardSnapshot.collected_by == "bek_audit",
                    CardSnapshot.match_status == "found",
                )
            )
            if row.fields.get("owner_claimed") is None
        ]
        blank_owner = blanks[0]
    create = _tasks(montebello_plan, "create_listing_candidate", "google")
    assert len(create) == 1
    assert create[0]["status"] == "unsupported"
    assert create[0]["completed"] is False
    assert create[0]["current"] == {"match_status": "not_found"}
    assert create[0]["editor"]["command"] is None
    assert google.fields["rating"] is None
    assert google.fields["reviews_count"] == 0
    assert google.field_sources["rating"]["csv_column"] == "google_рейтинг"
    assert google.field_sources["reviews_count"]["formula"] is False
    assert twogis.fields["whatsapp"] == "+7 776 111 25 85"
    prices = _tasks(badi_plan, "add_prices", "2gis")
    assert prices[0]["current"] == {"prices": "photo_only"}
    assert prices[0]["desired"] == {"prices": "company_price"}
    assert prices[0]["status"] == "blocked"
    assert prices[0]["completed"] is False
    assert prices[0]["editor"]["field"] == "services"
    replies = _tasks(cozy_plan, "prepare_review_replies", "2gis")
    assert replies[0]["current"] == {"replies_count": 0, "reviews_shown": 53}
    assert replies[0]["status"] == "manual"
    owner = _tasks(style_plan, "request_owner_verification", "google")
    assert owner[0]["current"] == {"owner_claimed": False}
    assert owner[0]["status"] == "manual"
    assert blank_owner is not None
    blank_plan = None
    with factory() as session:
        stored = session.scalar(select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == blank_owner.company_id))
        blank_plan = stored.plan
    assert _tasks(blank_plan, "request_owner_verification", "google") == []
    for plan in (montebello_plan, badi_plan, cozy_plan, style_plan):
        _assert_not_finished(plan)


def test_conflicting_matches_stay_for_a_person(factory, tmp_path: Path) -> None:
    path = tmp_path / "conflicts.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "карточки"
    headers = [
        "id_2gis",
        "название",
        "адрес_2гис",
        "дата_просмотра",
        "балл_итого",
        "балл_отзывы",
        "балл_цены",
        "балл_фото",
        "балл_контакты",
        "url_2gis",
        "номер_whatsapp",
        "яндекс_статус",
        "яндекс_рейтинг",
        "google_статус",
        "слой",
    ]
    sheet.append(headers)
    sheet.append(
        ["70000000000000001", "Дубль", "ул. Абая, 10", "2026-09-30", 40, 10, 10, 10, 10,
         "https://2gis.kz/almaty/firm/70000000000000001", "", "не найдена", "", "не найдена", "окраины"]
    )
    sheet.append(
        ["70000000000000002", "Дубль", "ул. Абая, 10", "2026-09-30", 41, 10, 10, 10, 11,
         "https://2gis.kz/almaty/firm/70000000000000002", "", "не найдена", "", "не найдена", "окраины"]
    )
    sheet.append(
        ["", "Дубль", "ул. Абая, 10", "2026-09-30", 42, 10, 10, 10, 12, "", "+7 701 111 22 33",
         "не найдена", "", "не найдена", "окраины"]
    )
    sheet.append(
        ["70000000000000003", "Чужой салон", "ул. Сатпаева, 20", "2026-09-30", 30, 5, 5, 5, 15,
         "https://2gis.kz/almaty/firm/70000000000000999", "+7 701 000 00 01", "не найдена", "", "не найдена", "окраины"]
    )
    sheet.append(
        ["", "Другой", "ул. Сатпаева, 20", "2026-09-30", 33, 5, 5, 5, 18, "", "+7 701 000 00 01",
         "не найдена", "", "не найдена", "окраины"]
    )
    formula_row = [
        "70000000000000004", "Формула", "ул. Толе би, 1", "2026-09-30", 20, 5, 5, 5, 5,
        "https://2gis.kz/almaty/firm/70000000000000004", "", "найдена", "", "не найдена", "окраины",
    ]
    sheet.append(formula_row)
    sheet.cell(sheet.max_row, headers.index("яндекс_рейтинг") + 1, "=1-0.1")
    sheet.append(
        ["70000000000000005", "Старая оценка", "ул. Гоголя, 2", "2026-09-30", 10, 1, 1, 1, 7,
         "https://2gis.kz/almaty/firm/70000000000000005", "", "не найдена", "", "не найдена", "формула"]
    )
    book.save(path)
    with factory() as session, session.begin():
        session.add(
            Company(
                id_2gis="70000000000000077",
                name="Телефон",
                city="Алматы",
                whatsapp="+7 701 000 00 01",
                source="seed",
            )
        )
        session.flush()
        seeded = session.scalar(select(Company).where(Company.id_2gis == "70000000000000077"))
        session.add(
            CardSnapshot(
                company_id=seeded.id,
                platform="2gis",
                match_status="found",
                fields={"address": "мкр Самал-2, 32", "phones": ["+7 701 000 00 01"]},
                field_sources={},
                purpose="prospect",
                collected_at=datetime(2026, 9, 1, tzinfo=ALMATY),
                collected_by="seed",
                status="collected",
            )
        )
        session.flush()
        stats = import_bek_audit(session, path)
        ids = set(session.scalars(select(Company.id_2gis)))
        formula = session.scalar(select(Company).where(Company.id_2gis == "70000000000000004"))
        yandex = session.scalar(
            select(CardSnapshot).where(CardSnapshot.company_id == formula.id, CardSnapshot.platform == "yandex")
        )
        tasks = session.scalars(select(HumanTask).where(HumanTask.reason == "bek_audit_ambiguous")).all()
    assert "70000000000000001" in ids
    assert "70000000000000002" in ids
    assert "70000000000000003" not in ids
    assert "70000000000000005" not in ids
    assert session_name_count(factory) == 2
    assert any(item.details["reason"] == "same name and address on two company ids" for item in tasks)
    assert any("more than one company" in (item.details.get("reason") or "") for item in tasks)
    assert any("does not match" in (item.details.get("reason") or "") for item in tasks)
    assert any("different address" in (item.details.get("reason") or "") for item in tasks)
    assert yandex.fields["rating"] is None
    assert yandex.field_sources["rating"]["formula"] is True
    assert any(item["reason"] == "unverified_layer" for item in stats.unmatched)
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(ChangeLog)) == 0
        assert session.scalar(select(func.count()).select_from(ChangeItem)) == 0


def session_name_count(factory) -> int:
    with factory() as session:
        return session.scalar(select(func.count()).select_from(Company).where(Company.name == "Дубль"))


def test_newer_verified_snapshot_is_not_overwritten(factory, tmp_path: Path) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    with factory() as session, session.begin():
        company = Company(id_2gis="70000001066660465", name="Montebello", city="Алматы", source="etalon_2026_09_30")
        session.add(company)
        session.flush()
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="2gis",
                match_status="found",
                url="https://2gis.kz/almaty/firm/70000001066660465",
                fields={"photos_count": 7, "name": "Montebello"},
                field_sources={"photos_count": {"csv_column": "kept"}},
                purpose="prospect",
                collected_at=datetime(2026, 10, 11, tzinfo=ALMATY),
                collected_by="manual_import",
                status="verified",
            )
        )
        stats = import_bek_audit(session, path)
        kept = session.scalar(
            select(CardSnapshot).where(
                CardSnapshot.company_id == company.id,
                CardSnapshot.platform == "2gis",
                CardSnapshot.status == "verified",
            )
        )
        bek = session.scalar(
            select(CardSnapshot).where(
                CardSnapshot.company_id == company.id,
                CardSnapshot.platform == "2gis",
                CardSnapshot.collected_by == "bek_audit",
            )
        )
    assert kept.fields["photos_count"] == 7
    assert bek is None
    assert "70000001066660465 2gis" in stats.skipped_newer
    assert stats.companies == 49


def test_newer_collected_snapshot_is_not_overwritten(factory, tmp_path: Path) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    with factory() as session, session.begin():
        company = Company(id_2gis="70000001066660465", name="Montebello", city="Алматы", source="etalon_2026_09_30")
        session.add(company)
        session.flush()
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="2gis",
                match_status="found",
                url="https://2gis.kz/almaty/firm/70000001066660465",
                fields={"photos_count": 7, "name": "Montebello"},
                field_sources={"photos_count": {"csv_column": "kept"}},
                purpose="prospect",
                collected_at=datetime(2026, 10, 11, tzinfo=ALMATY),
                collected_by="manual_import",
                status="collected",
            )
        )
        kept_id = company.id
        stats = import_bek_audit(session, path)
        rows = list(
            session.scalars(
                select(CardSnapshot).where(CardSnapshot.company_id == kept_id, CardSnapshot.platform == "2gis")
            ).all()
        )
    assert len(rows) == 1
    assert rows[0].status == "collected"
    assert rows[0].fields["photos_count"] == 7
    assert rows[0].collected_by == "manual_import"
    assert "70000001066660465 2gis" in stats.skipped_newer


def test_ready_only_when_prerequisites_are_met(factory, tmp_path: Path) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    with factory() as session, session.begin():
        import_bek_audit(session, path)
        company = _company(session, "70000001025141518")
        stored = session.scalar(select(ListingImprovementPlan).where(ListingImprovementPlan.company_id == company.id))
        assert _tasks(stored.plan, "add_prices", "2gis")[0]["status"] == "blocked"
        session.add(
            CompanyProfile(
                company_id=company.id,
                version=1,
                data={
                    "services": [
                        {"name": "Маникюр", "price": 8000, "source_ref": "msg:local", "needs_confirmation": False}
                    ]
                },
            )
        )
        access = session.get(CompanyPlatformId, (company.id, "2gis"))
        access.access_status = "verified"
        snapshots_before = session.scalar(select(func.count()).select_from(CardSnapshot))
        plan = company_plan(session, company.id)
        snapshots_after = session.scalar(select(func.count()).select_from(CardSnapshot))
    prices = _tasks(plan, "add_prices", "2gis")[0]
    assert prices["status"] == "ready_for_preparation"
    assert prices["completed"] is False
    assert prices["editor"]["command"] == "prepare-2gis-listing"
    assert _tasks(plan, "create_listing_candidate", "google")[0]["status"] == "unsupported"
    assert snapshots_after == snapshots_before
    assert plan["listing_modified"] is False


def test_import_does_not_modify_listings(factory, tmp_path: Path) -> None:
    source = (ROOT / "zhol" / "audit" / "bek.py").read_text(encoding="utf-8")
    plan_source = (ROOT / "zhol" / "audit" / "plan.py").read_text(encoding="utf-8")
    assert "playwright" not in source and "run_listing" not in source and "apply_changes" not in source
    assert "playwright" not in plan_source and "run_listing" not in plan_source
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    with factory() as session, session.begin():
        import_bek_audit(session, path)
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(ChangeLog)) == 0
        assert session.scalar(select(func.count()).select_from(ChangeItem)) == 0


def test_cli_prints_a_plan(factory, tmp_path: Path, capsys) -> None:
    path = _workbook(tmp_path / "karty-salony-almaty-2026-10.xlsx")
    assert main(["import-bek-audit", "--file", str(path)], factory=factory) == 0
    with factory() as session:
        company = _company(session, "70000001066660465")
    assert main(["bek-audit-plan", "--company-id", str(company.id)], factory=factory) == 0
    out = capsys.readouterr().out
    assert "лист карточки" in out
    assert "планы +50" in out
    plan = json.loads(out[out.index("{") :])
    assert plan["id_2gis"] == "70000001066660465"
    assert plan["listing_modified"] is False
    assert plan["score_total"] == 72
    assert plan["branch"]["company_id"] == company.id
