"""E9.3 admin workflow. Offline fixtures only; the main worker runs scout batches."""

from __future__ import annotations

import io
import shutil
import socket
import threading
from datetime import timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.orm.attributes import flag_modified

from zhol.api.main import app
from zhol.db.base import get_session_factory
from zhol.db.models import CollectionBatch, CollectionBatchItem, Company, CompanyProfile, ContentItem, Job, ListingApproval, ListingSubmission, Report
from zhol.jobs.queue import recover
from zhol.timeutil import now
from zhol.worker import serve

ROOT = Path(__file__).resolve().parents[1]
STUB = ROOT / "tests" / "golden" / "review-engine" / "config"
ORIGINAL = ROOT / "source" / "bek-audit-original.xlsx"
PAGES = ROOT / "tests" / "fixtures" / "cabinet"
TOKEN = "test-admin-token"
HEADERS = {"Authorization": f"Bearer {TOKEN}", "X-Zhol-Request": "admin"}
READ = {"Authorization": f"Bearer {TOKEN}"}
AUDIT_COMPANY = "70000001066660465"


@pytest.fixture
def client(migrated_engine, monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    config = tmp_path / "config"
    config.mkdir()
    shutil.copy(STUB / "models.yaml", config / "models.yaml")
    shutil.copy(STUB / "review_policy.yaml", config / "review_policy.yaml")
    shutil.copy(ROOT / "config" / "scoring.yaml", config / "scoring.yaml")
    monkeypatch.setenv("ZHOL_ROOT", str(ROOT))
    monkeypatch.setenv("ZHOL_CONFIG_DIR", str(config))
    monkeypatch.setenv("ZHOL_ADMIN_TOKEN", TOKEN)
    monkeypatch.setenv("ZHOL_ADMIN_FILES", str(tmp_path / "admin-files"))
    for name in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GOOGLE_API_KEY"):
        monkeypatch.setenv(name, "placeholder")
    factory = sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)
    stop = threading.Event()
    worker = threading.Thread(
        target=serve,
        args=(get_session_factory(), "main"),
        kwargs={"notifier": lambda _message: None, "stop": stop.is_set, "poll_seconds": 0.05},
        daemon=True,
    )
    worker.start()
    try:
        with TestClient(app) as test_client:
            yield test_client, factory
    finally:
        stop.set()
        worker.join(20)


def _offline(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(*_args, **_kwargs):
        raise AssertionError("network access attempted")

    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr("zhol.llm.providers.default_providers", lambda: (_ for _ in ()).throw(AssertionError("paid provider")))


def _wait_job(test_client: TestClient, job_id: int) -> dict:
    last: dict = {}
    for _ in range(400):
        response = test_client.get(f"/api/v1/admin/jobs/{job_id}", headers=READ)
        assert response.status_code == 200
        last = response.json()
        if last["status"] in {"done", "dead"}:
            return last
        threading.Event().wait(0.05)
    return last


def _order_spec(key: str) -> dict:
    facts = [
        {"field": "address", "value": "ул. Абая 10", "source_ref": "msg:m1", "approved": True},
        {"field": "phones", "value": ["+7 701 111 22 33"], "source_ref": "msg:m1", "approved": True},
        {"field": "hours", "value": "10:00-20:00", "source_ref": "msg:m1", "approved": True},
        {"field": "website", "value": "https://salon.example", "source_ref": "msg:m1", "approved": True},
        {"field": "description", "value": "Салон на Абая", "source_ref": "msg:m1", "approved": True},
        {"field": "category", "value": "Парикмахерская", "source_ref": "msg:m1", "approved": True},
        {"field": "services", "value": [{"name": "Стрижка", "price": 8000}], "source_ref": "msg:m1", "approved": True},
    ]
    platforms = [
        ("2GIS", "70000000000000081"),
        ("Yandex", "yandex-1"),
        ("Google", "google-1"),
    ]
    return {
        "key": key,
        "company": {"id_2gis": "70000000000000921", "name": "Salon Test", "city": "Алматы"},
        "platforms": [
            {"platform": name, "action": "update", "external_id": external_id, "owner_authorization": "verified"}
            for name, external_id in platforms
        ],
        "messages": [{"id": "m1", "text": "Салон на Абая. Сайт, телефон, часы, рубрика и стрижка."}],
        "facts": facts,
        "materials": [{"kind": "logo", "path": "materials/logo.png", "ref": "logo", "original_name": "logo.png"}],
        "not_applicable": ["videos"],
        "changes": [
            {
                "platform": name,
                "field": "phones",
                "after": ["+7 701 111 22 33"],
                "source_ref": "msg:m1",
                "nap_consent_ref": "msg:m1",
            }
            for name, _external_id in platforms
        ],
    }


def _listing(test_client: TestClient, order_id: int, action: str, package_id: int, version: int, platform: str, operation: str, page: bytes):
    return test_client.post(
        f"/api/v1/admin/orders/{order_id}/listings/{action}",
        headers=HEADERS,
        data={
            "package_id": str(package_id),
            "version": str(version),
            "platform": platform,
            "operation": operation,
            "mode": "offline",
        },
        files={"file": ("page.html", page, "text/html")},
    )


def test_offline_workflow_links_each_admin_step(client, monkeypatch: pytest.MonkeyPatch) -> None:
    _offline(monkeypatch)
    test_client, factory = client
    opened = test_client.get("/api/v1/admin/summary", headers=READ)
    assert opened.status_code == 200
    assert opened.json()["bek_connected"] is False
    anonymous = test_client.get("/api/v1/admin/summary")
    assert anonymous.status_code == 401

    started = test_client.post(
        "/api/v1/admin/scout/batches",
        json={
            "source_mode": "fixture",
            "city": "Алматы",
            "category": "Салон красоты",
            "platforms": ["2gis"],
            "limit": 1,
            "catalog": "catalog.yaml",
        },
        headers=HEADERS,
    )
    assert started.status_code == 200, started.text
    finished = _wait_job(test_client, started.json()["id"])
    assert finished["status"] == "done"
    assert finished["batch_id"]
    exported = test_client.post(f"/api/v1/admin/scout/batches/{finished['batch_id']}/export", headers=HEADERS)
    assert exported.status_code == 200, exported.text
    sheet = test_client.get(exported.json()["download"], headers=READ)
    assert sheet.status_code == 200
    book = load_workbook(io.BytesIO(sheet.content), read_only=True)
    found = {str(cell.value) for row in book.active.iter_rows() for cell in row}
    book.close()
    assert "70000001000000011" in found

    workbook = ORIGINAL.read_bytes()
    checked = test_client.post("/api/v1/admin/audits/validate", headers=HEADERS, files={"file": ("bek.xlsx", workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert checked.status_code == 200, checked.text
    assert checked.json()["imported"] is False
    imported = test_client.post("/api/v1/admin/audits/import", headers=HEADERS, files={"file": ("bek.xlsx", workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert imported.status_code == 200, imported.text
    assert imported.json()["imported"] is True
    assert imported.json()["bek_contacted"] is False
    with factory() as session:
        audited = session.scalar(select(Company).where(Company.id_2gis == AUDIT_COMPANY))
        assert audited is not None
        audited_id = audited.id
    audits = test_client.get(f"/api/v1/admin/audits?company_id={audited_id}", headers=READ)
    assert audits.status_code == 200
    assert audits.json()["total"] >= 1
    detail = test_client.get(f"/api/v1/admin/audits/{audits.json()['items'][0]['id']}", headers=READ)
    assert detail.status_code == 200
    assert "нет в Google" in detail.text
    company = test_client.get(f"/api/v1/admin/companies/{audited_id}", headers=READ)
    assert company.status_code == 200
    assert company.json()["company"]["id_2gis"] == AUDIT_COMPANY
    assert company.json()["snapshots"]

    created = test_client.post("/api/v1/admin/orders", json=_order_spec("e93-flow"), headers=HEADERS)
    assert created.status_code == 200, created.text
    order_id = created.json()["order_id"]
    readiness = created.json()["readiness"]
    assert readiness["missing_materials"]
    assert "photos" in readiness["missing_materials"]
    assert readiness["missing_owner_permissions"] == []
    assert created.json()["bek_contacted"] is False
    stored = test_client.get(f"/api/v1/admin/orders/{order_id}", headers=READ)
    assert stored.json()["readiness"]["missing_materials"] == readiness["missing_materials"]

    prepared = test_client.post(f"/api/v1/admin/orders/{order_id}/content/prepare", headers=HEADERS)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["provider"] == "offline"
    content = test_client.get(f"/api/v1/admin/content?order_id={order_id}", headers=READ)
    row = next(item for item in content.json()["items"] if item["status"] == "done")
    approved = test_client.post(f"/api/v1/admin/content/{row['id']}/approve", json={"version": row["version"]}, headers=HEADERS)
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "checked"
    assert approved.json()["publication"] == "not_published"

    dry = test_client.post(f"/api/v1/admin/orders/{order_id}/listings/dry-run", headers=HEADERS)
    assert dry.status_code == 200, dry.text
    assert dry.json()["published"] is False
    platforms = {item["platform"] for item in dry.json()["outcomes"]}
    assert platforms == {"2gis", "yandex", "google"}
    listed = test_client.get(f"/api/v1/admin/listings?order_id={order_id}", headers=READ)
    package = listed.json()["items"][0]
    by_platform = {item["platform"]: item for item in package["targets"]}
    assert set(by_platform) == {"2gis", "yandex", "google"}
    assert all(item["operation"] == "update" for item in by_platform.values())
    pages = {
        "2gis": (PAGES / "2gis-editor.html").read_bytes(),
        "yandex": (PAGES / "yandex.html").read_bytes(),
        "google": (PAGES / "google-editor.html").read_bytes(),
    }
    for name in ("2gis", "yandex", "google"):
        draft = _listing(test_client, order_id, "prepare", package["package_id"], package["version"], name, "update", pages[name])
        assert draft.status_code == 200, draft.text
        assert draft.json()["status"] == "prepared"
        assert draft.json()["published"] is False
        assert draft.json()["platform"] == name
    twogis = _listing(test_client, order_id, "approve", package["package_id"], package["version"], "2gis", "update", pages["2gis"])
    assert twogis.status_code == 200, twogis.text
    assert twogis.json()["status"] == "approved", twogis.text
    early = _listing(test_client, order_id, "simulate", package["package_id"], package["version"], "yandex", "update", pages["yandex"])
    assert early.status_code == 200, early.text
    assert early.json()["published"] is False
    assert early.json()["status"] != "approved"
    detail = test_client.get(f"/api/v1/admin/orders/{order_id}", headers=READ)
    other_spec = _order_spec("e93-other")
    other_spec.pop("company")
    other_spec["company_id"] = detail.json()["company_id"]
    other_spec["messages"][0]["id"] = "m2"
    for fact in other_spec["facts"]:
        fact["source_ref"] = "msg:m2"
    for change in other_spec["changes"]:
        change["source_ref"] = "msg:m2"
        change["nap_consent_ref"] = "msg:m2"
    other_order = test_client.post("/api/v1/admin/orders", json=other_spec, headers=HEADERS)
    assert other_order.status_code == 200, other_order.text
    foreign = _listing(
        test_client,
        other_order.json()["order_id"],
        "approve",
        package["package_id"],
        package["version"],
        "yandex",
        "update",
        pages["yandex"],
    )
    assert foreign.status_code == 404
    with factory() as session:
        approvals = session.scalars(select(ListingApproval).where(ListingApproval.order_id == order_id)).all()
        assert {row.platform for row in approvals if row.status == "active"} == {"2gis"}
    wrong_order = _listing(test_client, order_id + 999, "approve", package["package_id"], package["version"], "yandex", "update", pages["yandex"])
    assert wrong_order.status_code == 404
    crossed = _listing(test_client, order_id, "approve", package["package_id"], package["version"], "yandex", "create", pages["yandex"])
    assert crossed.status_code == 200, crossed.text
    assert crossed.json()["status"] != "approved"
    for name in ("yandex", "google"):
        accepted = _listing(test_client, order_id, "approve", package["package_id"], package["version"], name, "update", pages[name])
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["status"] == "approved"
        assert accepted.json()["platform"] == name
    for name in ("2gis", "yandex", "google"):
        simulated = _listing(test_client, order_id, "simulate", package["package_id"], package["version"], name, "update", pages[name])
        assert simulated.status_code == 200, simulated.text
        assert simulated.json()["published"] is False
        assert simulated.json()["live_publication"] is False
        assert simulated.json()["platform"] == name
        assert simulated.json()["status"] != "published"
    with factory() as session:
        rows = session.scalars(select(ListingSubmission).where(ListingSubmission.order_id == order_id)).all()
        assert {row.platform for row in rows} >= {"2gis", "yandex", "google"}
        assert all(row.status != "published" for row in rows)

    report = test_client.post(f"/api/v1/admin/orders/{order_id}/reports/completion", headers=HEADERS)
    assert report.status_code == 200, report.text
    assert report.json()["live_verified"] is False
    assert report.json()["order_id"] == order_id
    downloaded = test_client.get(report.json()["downloads"]["json"], headers=READ)
    assert downloaded.status_code == 200
    body = downloaded.json()
    assert body["order_id"] == order_id
    assert body["live_verified"] is False
    assert body.get("completed") is not True or body.get("outcome") != "stored_verified"
    html = test_client.get(report.json()["downloads"]["html"], headers=READ)
    assert html.status_code == 200
    assert "Живая проверка" in html.text
    hidden = test_client.get("/api/v1/admin/reports/999999/completion.json", headers=READ)
    assert hidden.status_code == 404
    with factory() as session, session.begin():
        stored_report = session.get(Report, report.json()["id"])
        assert stored_report is not None
        stored_report.order_id = other_order.json()["order_id"]
    stolen = test_client.get(report.json()["downloads"]["json"], headers=READ)
    assert stolen.status_code == 404
    assert b"70000000000000921" not in stolen.content


def test_worker_restart_resumes_the_same_batch(client, monkeypatch: pytest.MonkeyPatch) -> None:
    test_client, factory = client
    started = threading.Event()
    release = threading.Event()

    def slow(db_factory, item_id, _store, _cfg):
        started.set()
        assert release.wait(15)
        with db_factory() as session, session.begin():
            item = session.get(CollectionBatchItem, item_id)
            if item is not None:
                item.status = "completed"

    monkeypatch.setattr("zhol.agents.scout.batch._process_item", slow)
    try:
        created = test_client.post(
            "/api/v1/admin/scout/batches",
            json={
                "source_mode": "fixture",
                "city": "Алматы",
                "category": "Салон красоты",
                "platforms": ["2gis"],
                "limit": 2,
                "catalog": "catalog.yaml",
            },
            headers=HEADERS,
        )
        assert created.status_code == 200, created.text
        job_id = created.json()["id"]
        assert started.wait(10)
        again = test_client.post(
            "/api/v1/admin/scout/batches",
            json={
                "source_mode": "fixture",
                "city": "Алматы",
                "category": "Салон красоты",
                "platforms": ["2gis"],
                "limit": 2,
                "catalog": "catalog.yaml",
            },
            headers=HEADERS,
        )
        assert again.status_code == 200
        assert again.json()["duplicate"] is True
        assert again.json()["id"] == job_id
        with factory() as session, session.begin():
            job = session.get(Job, job_id)
            assert job is not None and job.status == "running"
            job.run_after = now() - timedelta(seconds=5)
        recover(factory, notifier=lambda _message: None)
        with factory() as session, session.begin():
            job = session.get(Job, job_id)
            assert job is not None and job.status == "failed"
            job.run_after = now()
    finally:
        release.set()
    finished = _wait_job(test_client, job_id)
    assert finished["status"] == "done", finished
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(CollectionBatch)) == 1
        batch = session.scalar(select(CollectionBatch))
        assert batch is not None and batch.status == "completed"
        completed = session.scalar(
            select(func.count()).select_from(CollectionBatchItem).where(
                CollectionBatchItem.batch_id == batch.id,
                CollectionBatchItem.status == "completed",
            )
        )
        assert completed == 2


def test_content_reprepare_invalidates_the_previous_version(client, monkeypatch: pytest.MonkeyPatch) -> None:
    _offline(monkeypatch)
    test_client, factory = client
    created = test_client.post("/api/v1/admin/orders", json=_order_spec("e93-version"), headers=HEADERS)
    assert created.status_code == 200, created.text
    order_id = created.json()["order_id"]
    assert test_client.post(f"/api/v1/admin/orders/{order_id}/content/prepare", headers=HEADERS).status_code == 200
    with factory() as session:
        row = session.scalars(select(ContentItem).where(ContentItem.order_id == order_id, ContentItem.status == "done")).first()
        assert row is not None
        content_id = row.id
        version = int(row.brief["version"])
    assert test_client.post(f"/api/v1/admin/content/{content_id}/approve", json={"version": version}, headers=HEADERS).status_code == 200
    with factory() as session, session.begin():
        from zhol.db.models import Order

        item = session.get(ContentItem, content_id)
        order = session.get(Order, item.order_id)
        profile = session.scalar(
            select(CompanyProfile).where(CompanyProfile.company_id == order.company_id).order_by(CompanyProfile.version.desc())
        )
        assert profile is not None
        data = dict(profile.data)
        hours = dict(data["hours"])
        hours["value"] = "09:00-18:00"
        data["hours"] = hours
        profile.data = data
        flag_modified(profile, "data")
    again = test_client.post(f"/api/v1/admin/orders/{order_id}/content/prepare", headers=HEADERS)
    assert again.status_code == 200, again.text
    with factory() as session:
        stored = session.get(ContentItem, content_id)
        assert stored is not None
        assert int(stored.brief["version"]) == version + 1
        assert stored.status == "done"
        assert stored.brief.get("approval") is None
    stale = test_client.post(f"/api/v1/admin/content/{content_id}/approve", json={"version": version}, headers=HEADERS)
    assert stale.status_code == 409
    current = test_client.post(f"/api/v1/admin/content/{content_id}/approve", json={"version": version + 1}, headers=HEADERS)
    assert current.status_code == 200, current.text
    assert current.json()["version"] == version + 1


def test_bad_page_does_not_prepare_the_other_platforms(client, monkeypatch: pytest.MonkeyPatch) -> None:
    _offline(monkeypatch)
    test_client, _factory = client
    created = test_client.post("/api/v1/admin/orders", json=_order_spec("e93-partial"), headers=HEADERS)
    assert created.status_code == 200, created.text
    order_id = created.json()["order_id"]
    dry = test_client.post(f"/api/v1/admin/orders/{order_id}/listings/dry-run", headers=HEADERS)
    assert dry.status_code == 200, dry.text
    listed = test_client.get(f"/api/v1/admin/listings?order_id={order_id}", headers=READ)
    package = listed.json()["items"][0]
    broken = _listing(
        test_client,
        order_id,
        "prepare",
        package["package_id"],
        package["version"],
        "google",
        "update",
        b"<html><body>not a cabinet</body></html>",
    )
    assert broken.status_code in {200, 400}
    if broken.status_code == 200:
        assert broken.json()["status"] != "prepared"
        assert broken.json()["published"] is False
    ready = _listing(
        test_client,
        order_id,
        "prepare",
        package["package_id"],
        package["version"],
        "2gis",
        "update",
        (PAGES / "2gis-editor.html").read_bytes(),
    )
    assert ready.status_code == 200, ready.text
    assert ready.json()["status"] == "prepared"
    assert ready.json()["platform"] == "2gis"
    refreshed = test_client.get(f"/api/v1/admin/listings?order_id={order_id}", headers=READ)
    states = {item["platform"]: item["preparation"] for item in refreshed.json()["items"][0]["targets"]}
    assert states["2gis"] == "prepared"
    assert states["google"] == "not_prepared"
    assert states["yandex"] == "not_prepared"
