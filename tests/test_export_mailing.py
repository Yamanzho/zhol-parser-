from __future__ import annotations

from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import sessionmaker

from zhol.db.models import CardSnapshot, Company
from zhol.etalon import import_etalon
from zhol.export.mailing import COLUMNS, build_row, dedupe_rows, load_categories, load_rows, write_xlsx
from zhol.timeutil import ALMATY

ROOT = Path(__file__).resolve().parents[1]
ETALON = ROOT / "source" / "etalon-50.csv"
WHEN = datetime(2026, 9, 30, tzinfo=ALMATY)


def _snap(category: str | None = None, **fields: object) -> SimpleNamespace:
    data = {"category": category}
    data.update(fields)
    return SimpleNamespace(
        fields=data,
        url="https://example.test/card",
        screenshot_path=None,
        status="verified",
        match_status="found",
        collected_at=WHEN,
        fail_reason=None,
        field_sources={},
    )


def _row(name: str, category: str | None, **kwargs: object) -> dict:
    table = load_categories()
    snapshots = {"2gis": _snap(category)}
    return build_row(
        company_id=str(kwargs.get("company_id") or name),
        name=name,
        city="Алматы",
        niche=kwargs.get("niche"),  # type: ignore[arg-type]
        source="test",
        whatsapp=None,
        instagram=None,
        snapshots=snapshots,
        services=kwargs.get("services"),  # type: ignore[arg-type]
        score=kwargs.get("score") if "score" in kwargs else 10,  # type: ignore[arg-type]
        discrepancies=[],
        categories=table,
    )


def test_columns_stay_in_order(tmp_path: Path) -> None:
    path = tmp_path / "mailing.xlsx"
    write_xlsx([_row("Nord", "Салон красоты", company_id="1")], path)
    sheet = load_workbook(path).active
    assert tuple(cell.value for cell in sheet[1]) == COLUMNS
    assert len(COLUMNS) == 40


def test_known_categories_normalize_and_unknown_stay_blank() -> None:
    salon = _row("Any name", "Салон красоты")
    gym = _row("Any name", "Фитнес-клуб")
    restaurant = _row("Any name", "Ресторан")
    medical = _row("Any name", "Стоматология")
    unknown = _row("Стоматология Улыбка", "Планетарий")
    ambiguous = _row("Montebello", None, niche="салон", services=None)
    assert (salon["Business Sector"], salon["Business Category"], salon["Business Subcategory"]) == ("Красота", "Салон красоты", None)
    assert (gym["Business Sector"], gym["Business Category"]) == ("Спорт", "Фитнес-клуб")
    assert (restaurant["Business Sector"], restaurant["Business Category"]) == ("Питание", "Ресторан")
    assert (medical["Business Sector"], medical["Business Category"]) == ("Медицина", "Стоматология")
    assert unknown["Business Sector"] is None and unknown["Business Category"] is None
    assert ambiguous["Business Sector"] is None
    assert unknown["Services / Activities"] is None
    assert salon["Services / Activities"] is None


def test_services_come_only_from_confirmed_profile_entries() -> None:
    row = _row(
        "Montebello",
        "Салон красоты",
        services=[
            {"name": "Стрижка", "needs_confirmation": False},
            {"name": "Окрашивание"},
            {"name": "Маникюр", "needs_confirmation": True},
        ],
    )
    assert row["Services / Activities"] == "Стрижка; Окрашивание"


def test_synonym_labels_share_one_mapping_and_disagreement_is_blank() -> None:
    table = load_categories()
    same = build_row(
        company_id="1",
        name="A",
        city=None,
        niche=None,
        source=None,
        whatsapp=None,
        instagram=None,
        snapshots={"2gis": _snap("Салон красоты"), "yandex": _snap("Beauty salon")},
        services=None,
        score=1,
        discrepancies=[],
        categories=table,
    )
    split = build_row(
        company_id="2",
        name="B",
        city=None,
        niche=None,
        source=None,
        whatsapp=None,
        instagram=None,
        snapshots={"2gis": _snap("Салон красоты"), "google": _snap("Ресторан")},
        services=None,
        score=1,
        discrepancies=[],
        categories=table,
    )
    assert same["Business Category"] == "Салон красоты"
    assert split["Business Sector"] is None and split["Business Category"] is None


def test_branch_rows_dedupe_by_company_and_keep_the_later_date() -> None:
    older = _row("Montebello", None, company_id="700")
    newer = _row("Montebello", None, company_id="700")
    other = _row("Montebello", None, company_id="701")
    older["Collection Date"] = "2026-09-30"
    newer["Collection Date"] = "2026-10-01"
    newer["Address"] = "другой адрес"
    other["Collection Date"] = "2026-09-30"
    rows = dedupe_rows([older, other, newer])
    assert [row["Company ID"] for row in rows] == ["700", "701"]
    assert rows[0]["Address"] == "другой адрес"


def test_null_stays_blank_zero_stays_zero_and_formulas_are_text(tmp_path: Path) -> None:
    table = load_categories()
    row = build_row(
        company_id="1",
        name="=1+1",
        city=None,
        niche=None,
        source="test",
        whatsapp="+7 700 000 00 00",
        instagram=None,
        snapshots={
            "2gis": _snap(None, rating=4.9, reviews_count=0, photos_count=None),
            "google": SimpleNamespace(
                fields={"reviews_count": 0},
                url=None,
                screenshot_path=None,
                status="verified",
                match_status="not_found",
                collected_at=WHEN,
                fail_reason=None,
                field_sources={},
            ),
        },
        services=None,
        score=72,
        discrepancies=[],
        categories=table,
    )
    path = tmp_path / "mailing.xlsx"
    assert write_xlsx([row], path) == 1
    sheet = load_workbook(path).active
    header = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, header["2GIS Review Count"]).value == 0
    assert sheet.cell(2, header["2GIS Photo Count"]).value is None
    assert sheet.cell(2, header["Google Review Count"]).value == 0
    assert sheet.cell(2, header["Google Rating"]).value is None
    assert sheet.cell(2, header["Google Photo Count"]).value is None
    assert sheet.cell(2, header["Company Name"]).value == "'=1+1"
    assert sheet.cell(2, header["WhatsApp"]).value == "'+7 700 000 00 00"
    assert sheet.cell(2, header["Overall Completeness Score"]).value == 72


def test_reference_export_uses_stored_snapshots(tmp_path: Path, migrated_engine: Engine) -> None:
    if not ETALON.is_file():
        pytest.skip("Файл source/etalon-50.csv не входит в публичный репозиторий.")
    factory = sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)
    with factory() as session, session.begin():
        import_etalon(session, ETALON)
    with factory() as session:
        rows = load_rows(session)
    path = tmp_path / "mailing.xlsx"
    assert write_xlsx(rows, path) == 50
    saved = {row["Company ID"]: row for row in rows}
    monte = saved["70000001066660465"]
    assert monte["Overall Completeness Score"] == 72
    assert monte["2GIS Review Count"] == 774
    assert monte["Google Review Count"] == 0
    assert monte["Google Rating"] is None
    assert monte["Business Sector"] is None
    assert monte["Services / Activities"] is None
    assert monte["Email"] is None
    assert monte["2GIS Description"] is None
    assert monte["Collection Date"] == "2026-09-30"
    assert monte["Verification Status"] == "verified"
    assert "google: not_found" in monte["Identified Problems"]
    company = None
    with factory() as session, session.begin():
        company = session.scalar(select(Company).where(Company.id_2gis == "70000001066660465"))
        assert company is not None
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="2gis",
                match_status="found",
                url="https://2gis.kz/almaty/firm/70000001066660465",
                fields={},
                field_sources={"_access": {"access": "robot_check", "outcome": "access_blocked"}},
                purpose="prospect",
                collected_at=datetime(2026, 10, 11, tzinfo=ALMATY),
                collected_by="test",
                status="failed",
                fail_reason="robot_check",
            )
        )
        company.dnc = False
    with factory() as session:
        again = load_rows(session)
    assert "70000001066660465" not in {row["Company ID"] for row in again}
    assert len(again) == 49
