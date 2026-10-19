"""Google names and source URLs stay on the Google snapshot.

These tests do not open a live map site and do not touch the application database.
"""

from __future__ import annotations

from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.google_card import apply_google_repair, assess_card, google_place_ready, repair_plan
from zhol.agents.scout.maps_selenium import MapPlace, _commit_place_locked, normalize_place
from zhol.agents.scout.platform_rows import list_platform_rows, write_platform_workbook
from zhol.db.models import CardSnapshot, CollectionBatch, CollectionBatchItem, Company
from zhol.timeutil import now

ROOT = Path(__file__).resolve().parents[1]
MAPS = "https://www.google.com/maps/place/data=!1sChIJstudio"


@pytest.fixture
def factory(migrated_engine):
    return sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)


def test_audit_evidence_has_no_google_name_or_maps_url() -> None:
    stored = assess_card(
        "google",
        {"name": None, "website": "starnails.kz", "rating": 4.6, "reviews_count": 58},
        None,
    )
    assert stored["name"] is None
    assert stored["source_url"] is None
    assert stored["completeness"] == "incomplete"
    assert stored["missing"] == ["name", "source_url"]
    assert google_place_ready("", None, "ChIJstudio") is False
    assert google_place_ready("Studio", "https://starnails.kz", "ChIJstudio") is False
    assert google_place_ready("Studio", MAPS, "ChIJstudio") is True
    local = "http://127.0.0.1:9/maps/place/data=!1sChIJstudio"
    assert google_place_ready("Studio", local, "ChIJstudio") is True
    assert assess_card("google", {"name": "Studio"}, local)["completeness"] == "complete"


def test_a_card_without_a_name_is_not_accepted() -> None:
    raw = {"name": "", "text": "", "hrefs": [], "address": "", "phones": [], "website": "starnails.kz"}
    assert (
        normalize_place(raw, platform="google", external_id="ChIJstudio", url=MAPS, city="Алматы", query="Салон")
        is None
    )
    kept = normalize_place(
        {**raw, "name": "Studio Google"},
        platform="google",
        external_id="ChIJstudio",
        url=MAPS,
        city="Алматы",
        query="Салон",
    )
    assert kept is not None
    assert kept.name == "Studio Google"
    assert kept.url == MAPS


def test_incomplete_google_rows_stay_empty_and_are_marked(factory, tmp_path: Path) -> None:
    when = now()
    with factory() as session, session.begin():
        company = Company(name="Стиль", city="Алматы", niche="салон", source="etalon")
        session.add(company)
        session.flush()
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="2gis",
                match_status="found",
                url="https://2gis.kz/almaty/firm/70000001054653678",
                fields={"name": "Стиль", "phones": ["+7 775 359 10 99"], "city": "Алматы"},
                field_sources={"name": {"csv_column": "название", "source": "bek_audit"}},
                purpose="prospect",
                collected_at=when,
                collected_by="bek_audit",
                status="collected",
            )
        )
        google = CardSnapshot(
            company_id=company.id,
            platform="google",
            match_status="found",
            url=None,
            fields={
                "name": None,
                "website": "starnails.kz",
                "rating": 4.6,
                "reviews_count": 58,
                "owner_claimed": True,
            },
            field_sources={
                "name": {"url": None, "csv_column": None, "source": "bek_audit"},
                "website": {"url": None, "csv_column": "google_сайт", "source": "bek_audit"},
                "rating": {"csv_column": "google_рейтинг", "source": "bek_audit"},
            },
            purpose="prospect",
            collected_at=when,
            collected_by="bek_audit",
            status="collected",
        )
        session.add(google)
        session.flush()
        google_id = google.id
        company_id = company.id
        before = (google.fields, google.url, google.collected_by)

    with factory() as session:
        row = session.get(CardSnapshot, google_id)
        assert row is not None
        plan = repair_plan(row)
        assert plan["action"] == "incomplete"
        assert plan["recoverable"] is False
        result = apply_google_repair(session, google_id)
        session.commit()
    assert result["action"] == "incomplete"
    with factory() as session:
        stored = session.get(CardSnapshot, google_id)
        assert stored is not None
        assert (stored.fields, stored.url, stored.collected_by) == before
        assert session.scalar(select(func.count()).select_from(CardSnapshot)) == 2
        listed = list_platform_rows(session, "google")
        item = listed["items"][0]
        assert item["company_id"] == company_id
        assert item["name"] is None
        assert item["source_url"] is None
        assert item["rating"] == 4.6
        assert item["completeness_label"] == "неполная"
        gis = list_platform_rows(session, "2gis")["items"][0]
        assert gis["name"] == "Стиль"
        assert gis["completeness_label"] == "полная"
        path = tmp_path / "platforms.xlsx"
        counts = write_platform_workbook(session, path)
    assert counts == {"2gis": 1, "google": 0}
    book = load_workbook(path)
    assert book.sheetnames == ["2ГИС", "Google Maps"]
    google_sheet = book["Google Maps"]
    headers = [cell.value for cell in google_sheet[1]]
    assert headers[-1] == "Полнота"
    assert "Адрес" in headers
    assert "Телефон" in headers
    assert list(google_sheet.iter_rows(min_row=2, values_only=True)) == []
    gis_rows = list(book["2ГИС"].iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in gis_rows] == ["Стиль"]
    assert "+7 775 359 10 99" not in {cell for row in book["Google Maps"].iter_rows(values_only=True) for cell in row}


def test_recorded_google_evidence_can_be_appended_without_rewriting_history(factory) -> None:
    when = now()
    with factory() as session, session.begin():
        company = Company(name="Общая", city="Алматы", source="etalon")
        session.add(company)
        session.flush()
        original = CardSnapshot(
            company_id=company.id,
            platform="google",
            match_status="found",
            url=None,
            fields={"name": None, "rating": 5},
            field_sources={
                "name": {"value": "Studio Google", "hyperlink": MAPS, "source": "google_html"},
            },
            purpose="prospect",
            collected_at=when,
            collected_by="saved_html",
            status="collected",
        )
        session.add(original)
        session.flush()
        original_id = original.id
        original_fields = dict(original.fields)

    with factory() as session, session.begin():
        result = apply_google_repair(session, original_id)
        assert result["action"] == "recovered"
        assert result["name"] == "Studio Google"
        assert result["source_url"] == MAPS
    with factory() as session:
        old = session.get(CardSnapshot, original_id)
        assert old is not None
        assert old.fields == original_fields
        assert old.url is None
        created = session.get(CardSnapshot, result["snapshot_id"])
        assert created is not None
        assert created.id != original_id
        assert created.fields["name"] == "Studio Google"
        assert created.url == MAPS
        shown = list_platform_rows(session, "google")["items"][0]
        assert shown["name"] == "Studio Google"
        assert shown["source_url"] == MAPS
        assert shown["completeness_label"] == "полная"


def test_an_incomplete_live_card_is_rejected_before_it_is_stored(factory) -> None:
    with factory() as session, session.begin():
        batch = CollectionBatch(
            city="Алматы",
            category="Салон красоты",
            platforms=["google"],
            item_limit=1,
            source_mode="browser_google",
            status="running",
            out_path="/tmp/zhol-google-incomplete.xlsx",
            catalog_path="browser_google",
            progress={},
        )
        session.add(batch)
        session.flush()
        outcome = _commit_place_locked(session, batch.id, _blank_place(), 0.0)
        assert outcome == "rejected"
        assert session.scalar(select(func.count()).select_from(Company)) == 0
        assert session.scalar(select(func.count()).select_from(CardSnapshot)) == 0
        item = session.scalar(select(CollectionBatchItem))
        assert item is not None
        assert item.status == "blocked"
        assert item.detail["classification"] == "incomplete_card"


def test_newer_incomplete_google_snapshot_does_not_hide_the_complete_one(factory, tmp_path: Path) -> None:
    earlier = now() - timedelta(days=1)
    later = now()
    maps = "https://www.google.com/maps/place/Studio/data=!1sChIJcomplete"
    with factory() as session, session.begin():
        company = Company(name="Studio", city="Алматы", niche="салон", source="test")
        session.add(company)
        session.flush()
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="google",
                match_status="found",
                url=maps,
                fields={"name": "Studio Google", "address": "Abay 1", "phones": ["+7 701 000 00 01"]},
                field_sources={},
                purpose="prospect",
                collected_at=earlier,
                collected_by="live",
                status="collected",
            )
        )
        session.add(
            CardSnapshot(
                company_id=company.id,
                platform="google",
                match_status="found",
                url=None,
                fields={"name": None, "website": "starnails.kz", "rating": 4.6},
                field_sources={},
                purpose="prospect",
                collected_at=later,
                collected_by="bek_audit",
                status="collected",
            )
        )
        company_id = company.id
    with factory() as session:
        page = list_platform_rows(session, "google")
        assert page["total"] == 1
        assert page["items"][0]["name"] == "Studio Google"
        assert page["items"][0]["source_url"] == maps
        history = list(session.scalars(select(CardSnapshot).where(CardSnapshot.company_id == company_id)).all())
        assert len(history) == 2
        assert {row.collected_by for row in history} == {"live", "bek_audit"}
        write_platform_workbook(session, tmp_path / "platforms.xlsx")
    book = load_workbook(tmp_path / "platforms.xlsx")
    names = [row[0] for row in book["Google Maps"].iter_rows(min_row=2, values_only=True)]
    link_at = [cell.value for cell in book["Google Maps"][1]].index("Ссылка на карточку")
    links = [row[link_at] for row in book["Google Maps"].iter_rows(min_row=2, values_only=True)]
    assert names == ["Studio Google"]
    assert links == [maps]


def _blank_place() -> MapPlace:
    return MapPlace(
        platform="google",
        external_id="ChIJblank",
        name="",
        city="Алматы",
        shown_city=None,
        address=None,
        address_reason="not_shown",
        phones=(),
        phone_display=(),
        phones_reason="not_shown",
        whatsapp=None,
        whatsapp_reason="not_shown",
        instagram=None,
        instagram_reason="not_shown",
        website="starnails.kz",
        website_reason=None,
        rating=4.6,
        rating_reason=None,
        review_count=58,
        review_reason=None,
        category=None,
        category_reason="not_shown",
        url="",
        query="Салон красоты",
        collected_at=now().isoformat(),
    )
