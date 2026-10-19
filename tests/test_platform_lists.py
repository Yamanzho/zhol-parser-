"""Separate 2GIS and Google lists. These tests do not open a live map site."""

from __future__ import annotations

import os
import subprocess
import sys
from datetime import timedelta
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from zhol.agents.scout.browser_2gis import GRANT_LOCAL
from zhol.agents.scout.platform_rows import list_platform_rows, write_platform_workbook
from zhol.api.main import app
from zhol.browser.profile_gate import exclusive_chrome_profile
from zhol.db.models import CardSnapshot, CollectionBatch, Company, CompanyPlatformId, Job
from zhol.jobs.queue import JobDeferred, enqueue, run_one
from zhol.timeutil import now

ROOT = Path(__file__).resolve().parents[1]
STUB = ROOT / "tests" / "golden" / "review-engine" / "config"
TOKEN = "test-admin-token"
HEADERS = {"Authorization": f"Bearer {TOKEN}", "X-Zhol-Request": "admin"}


@pytest.fixture
def factory(migrated_engine, monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    monkeypatch.setenv("ZHOL_ROOT", str(ROOT))
    monkeypatch.setenv("ZHOL_CONFIG_DIR", str(STUB))
    monkeypatch.setenv("ZHOL_ADMIN_TOKEN", TOKEN)
    monkeypatch.setenv("ZHOL_ADMIN_FILES", str(tmp_path / "admin-files"))
    monkeypatch.setenv("ZHOL_BROWSER_ENGINE", "1")
    monkeypatch.setenv("ZHOL_2GIS_BROWSER_GRANT", GRANT_LOCAL)
    monkeypatch.setenv("ZHOL_GOOGLE_BROWSER_GRANT", GRANT_LOCAL)
    monkeypatch.delenv("ZHOL_YANDEX_BROWSER_GRANT", raising=False)
    monkeypatch.delenv("ZHOL_BROWSER_BASE_URL", raising=False)
    monkeypatch.delenv("ZHOL_BROWSER_HTML_ROOT", raising=False)
    monkeypatch.delenv("DGIS_CATALOG_KEY", raising=False)
    for name in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GOOGLE_API_KEY"):
        monkeypatch.setenv(name, "placeholder")
    return sessionmaker(bind=migrated_engine, expire_on_commit=False, future=True)


def _company(session, name: str) -> Company:
    company = Company(name=name, city="Алматы", niche="Салон красоты", source="saved_pages")
    session.add(company)
    session.flush()
    return company


def _snapshot(session, company: Company, platform: str, fields: dict, url: str, when) -> CardSnapshot:
    row = CardSnapshot(
        company_id=company.id,
        platform=platform,
        match_status="found",
        url=url,
        fields=fields,
        field_sources={},
        purpose="prospect",
        collected_at=when,
        collected_by="test",
        status="collected",
    )
    session.add(row)
    session.flush()
    if session.get(CompanyPlatformId, (company.id, platform)) is None:
        session.add(
            CompanyPlatformId(
                company_id=company.id,
                platform=platform,
                external_id=f"{platform}-{company.id}",
                public_url=url,
            )
        )
    return row


def _seed(factory) -> dict:
    earlier = now() - timedelta(days=2)
    later = now() - timedelta(days=1)
    with factory() as session, session.begin():
        shared = _company(session, "Общая студия")
        only_gis = _company(session, "Только 2ГИС")
        only_google = _company(session, "Только Google")
        _snapshot(
            session,
            shared,
            "2gis",
            {"name": "Старое имя 2ГИС", "phones": ["+7 701 000 00 01"], "city": "Алматы", "address": "старый адрес"},
            "https://2gis.kz/old",
            earlier,
        )
        gis = _snapshot(
            session,
            shared,
            "2gis",
            {
                "name": "Студия 2ГИС",
                "phones": ["+7 701 111 11 11"],
                "whatsapp": "https://wa.me/77011111111",
                "instagram": "https://instagram.com/studio_gis",
                "city": "Алматы",
                "address": "улица Абая, 1",
                "category": "Салон красоты",
                "rating": 4.6,
                "reviews_count": 12,
            },
            "https://2gis.kz/firm/shared",
            later,
        )
        google = _snapshot(
            session,
            shared,
            "google",
            {
                "name": "Studio Google",
                "phones": ["+7 701 222 22 22"],
                "city": "Алматы",
                "address": "Abay 1",
                "category": "Beauty salon",
                "rating": 4.2,
                "reviews_count": 8,
            },
            "https://www.google.com/maps/place/shared",
            later,
        )
        _snapshot(
            session,
            only_gis,
            "2gis",
            {"name": "Только 2ГИС", "phones": ["+7 701 333 33 33"], "city": "Алматы", "address": "улица Достык, 2"},
            "https://2gis.kz/firm/only",
            later,
        )
        _snapshot(
            session,
            only_google,
            "google",
            {"name": "Google Only", "phones": ["+7 701 444 44 44"], "city": "Алматы", "address": "Dostyk 3"},
            "https://www.google.com/maps/place/only",
            later,
        )
        _snapshot(
            session,
            only_gis,
            "yandex",
            {"name": "Не показывать в этих вкладках", "phones": ["+7 701 555 55 55"], "city": "Алматы"},
            "https://yandex.kz/maps/org/hidden",
            later,
        )
        hidden = CardSnapshot(
            company_id=only_google.id,
            platform="google",
            match_status="not_found",
            url="https://www.google.com/maps/place/missing",
            fields={"name": "Не найдена"},
            field_sources={},
            purpose="prospect",
            collected_at=later,
            collected_by="test",
            status="collected",
        )
        session.add(hidden)
        batch = CollectionBatch(
            city="Алматы",
            category="Салон красоты",
            platforms=["2gis"],
            item_limit=20,
            source_mode="browser_2gis",
            status="completed",
            out_path="/tmp/zhol-existing-batch.xlsx",
            catalog_path="browser_2gis",
        )
        session.add(batch)
        session.flush()
        return {
            "shared": shared.id,
            "gis_snapshot": gis.id,
            "google_snapshot": google.id,
            "batch": batch.id,
            "gis_fields": dict(gis.fields),
            "google_fields": dict(google.fields),
        }


def _collect(client: TestClient, modes: list[str], limit: int = 100) -> object:
    platforms = ["google" if mode == "browser_google" else "2gis" for mode in modes]
    return client.post(
        "/api/v1/admin/scout/collect",
        headers=HEADERS,
        json={
            "city": "Алматы",
            "queries": ["Салон красоты"],
            "platforms": platforms,
            "limit": limit,
            "source_mode": "fixture",
            "source_modes": modes,
        },
    )


def test_each_source_and_both_enqueue_separate_batches(factory) -> None:
    with TestClient(app) as client:
        gis = _collect(client, ["browser_2gis"])
        google = _collect(client, ["browser_google"])
        both = _collect(client, ["browser_2gis", "browser_google"])
    assert gis.status_code == 200
    assert google.status_code == 200
    assert both.status_code == 200
    assert [job["source_mode"] for job in gis.json()["jobs"]] == ["browser_2gis"]
    assert [job["source_mode"] for job in google.json()["jobs"]] == ["browser_google"]
    paired = both.json()
    assert paired["separate_batches"] is True
    assert paired["limit"] == 100
    assert [job["source_mode"] for job in paired["jobs"]] == ["browser_2gis", "browser_google"]
    assert [job["duplicate"] for job in paired["jobs"]] == [True, True]
    with factory() as session:
        jobs = list(session.scalars(select(Job).order_by(Job.id)).all())
        assert [job.payload["source_mode"] for job in jobs] == ["browser_2gis", "browser_google"]
        assert [job.payload["limit"] for job in jobs] == [100, 100]
        assert [job.payload["platforms"] for job in jobs] == [["2gis"], ["google"]]
        assert {job.queue for job in jobs} == {"collector"}
        assert session.scalar(select(func.count()).select_from(Company)) == 0
        assert session.scalar(select(func.count()).select_from(CollectionBatch)) == 0


def test_a_closed_source_does_not_cancel_the_other(factory, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("ZHOL_GOOGLE_BROWSER_GRANT", raising=False)
    with TestClient(app) as client:
        response = _collect(client, ["browser_2gis", "browser_google"], limit=10)
    assert response.status_code == 200
    body = response.json()
    assert [job["source_mode"] for job in body["jobs"]] == ["browser_2gis"]
    assert body["errors"][0]["source_mode"] == "browser_google"
    with factory() as session:
        jobs = list(session.scalars(select(Job)).all())
        assert len(jobs) == 1
        assert jobs[0].payload["source_mode"] == "browser_2gis"
        assert jobs[0].payload["limit"] == 10


def test_tables_keep_platform_values_and_one_company(factory, tmp_path: Path) -> None:
    seeded = _seed(factory)
    with factory() as session:
        before_companies = session.scalar(select(func.count()).select_from(Company))
        before_snapshots = [
            (row.id, row.company_id, row.platform, row.fields, row.status)
            for row in session.scalars(select(CardSnapshot).order_by(CardSnapshot.id)).all()
        ]
        batch = session.get(CollectionBatch, seeded["batch"])
        assert batch is not None
        before_batch = (batch.status, batch.city, list(batch.platforms), batch.out_path)
        gis = list_platform_rows(session, "2gis", q="", limit=50, offset=0)
        google = list_platform_rows(session, "google", q="222", limit=50, offset=0)
        page = list_platform_rows(session, "2gis", q="", limit=1, offset=1)
        counts = write_platform_workbook(session, tmp_path / "platforms.xlsx")
    gis_ids = [row["company_id"] for row in gis["items"]]
    assert seeded["shared"] in gis_ids
    shared_gis = next(row for row in gis["items"] if row["company_id"] == seeded["shared"])
    assert shared_gis["name"] == "Студия 2ГИС"
    assert shared_gis["phone"] == "+7 701 111 11 11"
    assert shared_gis["address"] == "улица Абая, 1"
    assert shared_gis["rating"] == 4.6
    assert shared_gis["reviews_count"] == 12
    assert shared_gis["source_url"] == "https://2gis.kz/firm/shared"
    assert shared_gis["category"] == "Салон красоты"
    assert shared_gis["search_query"] is None
    only_row = next(row for row in gis["items"] if row["name"] == "Только 2ГИС")
    assert only_row["category"] == "Не указана"
    assert "Studio Google" not in {row["name"] for row in gis["items"]}
    assert "Не показывать в этих вкладках" not in {row["name"] for row in gis["items"]}
    assert gis["total"] == 2
    assert len(gis_ids) == len(set(gis_ids))
    assert google["total"] == 1
    assert google["items"][0]["company_id"] == seeded["shared"]
    assert google["items"][0]["name"] == "Studio Google"
    assert google["items"][0]["phone"] == "+7 701 222 22 22"
    assert google["items"][0]["whatsapp"] is None
    assert page["total"] == 2
    assert len(page["items"]) == 1
    book = load_workbook(tmp_path / "platforms.xlsx")
    assert book.sheetnames == ["2ГИС", "Google Maps"]
    gis_sheet = book["2ГИС"]
    google_sheet = book["Google Maps"]
    assert gis_sheet.freeze_panes == "A2"
    assert gis_sheet.auto_filter.ref.startswith("A1:")
    assert [cell.value for cell in gis_sheet[1]][:6] == [
        "Название",
        "Категория",
        "Поисковый запрос",
        "Источник",
        "Город",
        "Регион",
    ]
    gis_names = [row[0] for row in gis_sheet.iter_rows(min_row=2, values_only=True)]
    google_names = [row[0] for row in google_sheet.iter_rows(min_row=2, values_only=True)]
    assert "Студия 2ГИС" in gis_names
    assert "Studio Google" not in gis_names
    assert "Studio Google" in google_names
    assert "Студия 2ГИС" not in google_names
    assert "Не найдена" not in google_names
    assert counts == {"2gis": 2, "google": 2}
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(Company)) == before_companies
        after_snapshots = [
            (row.id, row.company_id, row.platform, row.fields, row.status)
            for row in session.scalars(select(CardSnapshot).order_by(CardSnapshot.id)).all()
        ]
        assert after_snapshots == before_snapshots
        batch = session.get(CollectionBatch, seeded["batch"])
        assert batch is not None
        assert (batch.status, batch.city, list(batch.platforms), batch.out_path) == before_batch
        stored = session.get(CardSnapshot, seeded["gis_snapshot"])
        assert stored is not None and stored.fields == seeded["gis_fields"]


def test_the_api_lists_one_platform_and_downloads_both_sheets(factory) -> None:
    _seed(factory)
    with TestClient(app) as client:
        listed = client.get("/api/v1/admin/scout/platforms/2gis?limit=50&offset=0", headers=HEADERS)
        missing = client.get("/api/v1/admin/scout/platforms/yandex", headers=HEADERS)
        prepared = client.post("/api/v1/admin/scout/platform-export", headers=HEADERS)
        downloaded = client.get("/api/v1/admin/scout/platform-export.xlsx", headers=HEADERS)
    assert listed.status_code == 200
    assert listed.json()["label"] == "2ГИС"
    assert listed.json()["total"] == 2
    assert missing.status_code == 404
    assert prepared.status_code == 200
    assert prepared.json()["sheets"] == {"2gis": 2, "google": 2}
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_download_replaces_a_stale_workbook_and_prefers_the_complete_google_card(factory, tmp_path: Path) -> None:
    earlier = now() - timedelta(days=3)
    later = now() - timedelta(hours=1)
    with factory() as session, session.begin():
        company = _company(session, "Живая карточка")
        _snapshot(
            session,
            company,
            "google",
            {
                "name": "Google Live",
                "phones": ["+7 701 777 77 77"],
                "city": "Алматы",
                "address": "улица Абая, 7",
                "category": "Салон красоты",
                "rating": 4.8,
                "reviews_count": 21,
            },
            "https://www.google.com/maps/place/live",
            earlier,
        )
        audit = CardSnapshot(
            company_id=company.id,
            platform="google",
            match_status="found",
            url=None,
            fields={"rating": 4.1, "reviews_count": 3},
            field_sources={},
            purpose="prospect",
            collected_at=later,
            collected_by="bek_audit",
            status="collected",
        )
        session.add(audit)
        empty = _company(session, "Только аудит")
        session.add(
            CardSnapshot(
                company_id=empty.id,
                platform="google",
                match_status="found",
                url=None,
                fields={"rating": 3.2, "reviews_count": 1},
                field_sources={},
                purpose="prospect",
                collected_at=later,
                collected_by="manual_import",
                status="collected",
            )
        )
        before = session.scalar(select(func.count()).select_from(CardSnapshot))
    folder = tmp_path / "admin-files" / "exports"
    folder.mkdir(parents=True)
    stale = Workbook()
    stale_sheet = stale.active
    stale_sheet.title = "Google Maps"
    stale_sheet.append(["Номер компании", "Название", "Телефон", "WhatsApp", "Instagram", "Город", "Адрес", "Категория", "Ссылка", "Рейтинг", "Отзывы", "Дата сбора"])
    stale_sheet.append([1, None, None, None, None, None, None, None, None, 4.1, 3, "2026-10-01"])
    stale.save(folder / "platforms.xlsx")
    page = (ROOT / "web" / "src" / "collect.tsx").read_text(encoding="utf-8")
    main, extra = page.split("<details", 1)
    assert "downloadPlatformWorkbook()" in main
    assert "/api/v1/admin/scout/export.xlsx" not in main
    assert "/api/v1/admin/scout/export.xlsx" in extra
    companies = (ROOT / "web" / "src" / "companies.tsx").read_text(encoding="utf-8")
    assert "downloadPlatformWorkbook()" in companies
    with TestClient(app) as client:
        downloaded = client.get("/api/v1/admin/scout/platform-export.xlsx", headers=HEADERS)
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.headers["cache-control"] == "no-store"
    book = load_workbook(BytesIO(downloaded.content))
    assert book.sheetnames == ["2ГИС", "Google Maps"]
    headers = [cell.value for cell in book["Google Maps"][1]]
    assert headers[:6] == ["Название", "Категория", "Поисковый запрос", "Источник", "Город", "Регион"]
    assert "Ссылка на карточку" in headers
    assert "Полнота" in headers
    rows = list(book["Google Maps"].iter_rows(min_row=2, values_only=True))
    by_name = {row[0]: row for row in rows}
    live = by_name["Google Live"]
    assert live[headers.index("Категория")] == "Салон красоты"
    assert live[headers.index("Адрес")] == "улица Абая, 7"
    assert live[headers.index("Город")] == "Алматы"
    assert live[headers.index("Ссылка на карточку")] == "https://www.google.com/maps/place/live"
    assert live[headers.index("Телефон")] == "'+7 701 777 77 77"
    assert live[headers.index("Полнота")] == "полная"
    assert "Студия 2ГИС" not in by_name
    assert "Только аудит" not in by_name
    assert all(row[headers.index("Полнота")] == "полная" for row in rows)
    assert len(rows) == 1
    stored = load_workbook(folder / "platforms.xlsx")
    assert [cell.value for cell in stored["Google Maps"][1]] == headers
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(CardSnapshot)) == before


def _hold_profile(profile: Path) -> subprocess.Popen[str]:
    script = (
        "import os, time\n"
        "from zhol.browser.profile_gate import exclusive_chrome_profile\n"
        "with exclusive_chrome_profile():\n"
        "    print('held', flush=True)\n"
        "    time.sleep(30)\n"
    )
    env = os.environ.copy()
    env["ZHOL_CHROME_PROFILE"] = str(profile)
    env["PYTHONPATH"] = str(ROOT) + os.pathsep + env.get("PYTHONPATH", "")
    process = subprocess.Popen(
        [sys.executable, "-c", script],
        cwd=str(ROOT),
        env=env,
        stdout=subprocess.PIPE,
        text=True,
    )
    assert process.stdout is not None
    line = process.stdout.readline()
    if "held" not in line:
        process.kill()
        raise AssertionError(line or "profile lock did not start")
    return process


def test_a_busy_profile_waits_and_a_failed_source_leaves_the_other_queued(factory, monkeypatch, tmp_path: Path) -> None:
    profile = tmp_path / "chrome-profile"
    monkeypatch.setenv("ZHOL_CHROME_PROFILE", str(profile))
    holder = _hold_profile(profile)
    try:
        with pytest.raises(JobDeferred):
            with exclusive_chrome_profile():
                pass
    finally:
        holder.kill()
        holder.wait(timeout=5)
    with factory() as session, session.begin():
        first = enqueue(
            session,
            "admin.scout_batch",
            {"source_mode": "browser_2gis", "city": "Алматы", "category": "Салон красоты", "limit": 100},
            queue="collector",
        )
        second = enqueue(
            session,
            "admin.scout_batch",
            {"source_mode": "browser_google", "city": "Алматы", "category": "Салон красоты", "limit": 100},
            queue="collector",
        )
    seen: list[str] = []

    def fail_gis(payload: dict) -> None:
        if payload["source_mode"] == "browser_2gis":
            raise RuntimeError("2gis stopped")
        seen.append(payload["source_mode"])

    assert run_one(factory, "collector", {"admin.scout_batch": fail_gis}, notifier=lambda _message: None) == first
    assert run_one(factory, "collector", {"admin.scout_batch": fail_gis}, notifier=lambda _message: None) == second
    assert seen == ["browser_google"]
    with factory() as session:
        gis_job = session.get(Job, first)
        google_job = session.get(Job, second)
        assert gis_job is not None and gis_job.status == "failed"
        assert google_job is not None and google_job.status == "done"
        waiting = enqueue(
            session,
            "admin.scout_batch",
            {"source_mode": "browser_google", "wait": True},
            queue="collector",
            dedup_key="platform-list-wait",
        )
        session.commit()
    holder = _hold_profile(profile)
    try:

        def wait_for_profile(_payload: dict) -> None:
            with exclusive_chrome_profile():
                pass

        assert run_one(factory, "collector", {"admin.scout_batch": wait_for_profile}, notifier=lambda _message: None) == waiting
    finally:
        holder.kill()
        holder.wait(timeout=5)
    with factory() as session:
        job = session.get(Job, waiting)
        assert job is not None
        assert job.status == "queued"
        assert job.attempts == 0
        assert job.run_after > now()


@pytest.mark.parametrize("mode", ["browser_yandex", "browser_google", "browser_2gis"])
def test_a_locked_profile_defers_every_browser_collector(mode: str, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    import parser_service

    from zhol.api.admin_jobs import _handle
    from zhol.browser.discovery.types import RunControl

    profile = tmp_path / "chrome-profile"
    monkeypatch.setenv("ZHOL_CHROME_PROFILE", str(profile))
    released: list[object] = []
    monkeypatch.setattr(parser_service, "release_browser", lambda *args, **kwargs: released.append(args))
    monkeypatch.setattr("zhol.agents.scout.browser_yandex.collection_ready", lambda: True)
    monkeypatch.setattr("zhol.agents.scout.browser_google.collection_ready", lambda: True)
    monkeypatch.setattr("zhol.agents.scout.browser_2gis.browser_collection_authorized", lambda: True)

    def started(*_args: object, **_kwargs: object) -> None:
        parser_service.release_browser(profile)
        raise AssertionError("collector started while the profile lock was held")

    monkeypatch.setattr("zhol.api.admin_jobs._run_maps", started)
    monkeypatch.setattr("zhol.api.admin_jobs._run_browser", started)
    holder = _hold_profile(profile)
    try:
        with pytest.raises(JobDeferred):
            _handle(None, {"source_mode": mode, "platforms": [mode.removeprefix("browser_")]}, RunControl(), {}, 1, 1)
    finally:
        holder.kill()
        holder.wait(timeout=5)
    assert released == []
