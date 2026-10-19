"""Read-only lists of collected 2GIS and Google cards.

Each row comes from the latest snapshot of that platform. A shared company
keeps one identity and is not copied. A business category comes only from that
card. The search phrase is a separate column and is never written back.
"""

from __future__ import annotations

import os
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import String, and_, case, cast, func, or_, select
from sqlalchemy.orm import Session

from zhol.agents.scout.google_card import assess_card
from zhol.agents.scout.localities import region_of
from zhol.db.models import CardSnapshot, CollectionBatch

PLATFORMS = ("2gis", "google")
LABELS = {"2gis": "2ГИС", "google": "Google Maps"}
NOT_STATED = "Не указана"
UNCONFIRMED_CITY = "Город не подтверждён"
COLUMNS = (
    ("name", "Название"),
    ("category", "Категория"),
    ("search_query", "Поисковый запрос"),
    ("source", "Источник"),
    ("city", "Город"),
    ("region", "Регион"),
    ("address", "Адрес"),
    ("phone", "Телефон"),
    ("whatsapp", "WhatsApp"),
    ("instagram", "Instagram"),
    ("source_url", "Ссылка на карточку"),
    ("rating", "Рейтинг"),
    ("reviews_count", "Отзывы"),
    ("collected_at", "Дата сбора"),
    ("completeness_label", "Полнота"),
)
_TEXT_KEYS = ("name", "whatsapp", "instagram", "city", "address")
_QUERY_PREFIXES = frozenset({"browser_2gis", "browser_google", "browser_yandex", "live_2gis", "batch", "discovery"})
_REPO = Path(__file__).resolve().parents[3]


def platform_label(platform: str) -> str:
    if platform not in LABELS:
        raise ValueError("unknown platform")
    return LABELS[platform]


def list_platform_rows(
    session: Session,
    platform: str,
    *,
    q: str = "",
    limit: int = 50,
    offset: int = 0,
) -> dict[str, Any]:
    platform_label(platform)
    latest = _latest_ids(platform).subquery()
    stmt = select(CardSnapshot).where(CardSnapshot.id.in_(select(latest.c.id)))
    term = " ".join(q.split())
    if term:
        needle = f"%{_escape(term)}%"
        stmt = stmt.where(
            or_(
                cast(CardSnapshot.fields, String).ilike(needle, escape="\\"),
                CardSnapshot.url.ilike(needle, escape="\\"),
            )
        )
    total = int(session.scalar(select(func.count()).select_from(stmt.subquery())) or 0)
    rows = session.scalars(
        stmt.order_by(CardSnapshot.collected_at.desc(), CardSnapshot.id.desc()).limit(limit).offset(offset)
    ).all()
    queries = _batch_queries(session, rows)
    return {
        "platform": platform,
        "label": LABELS[platform],
        "total": total,
        "limit": limit,
        "offset": offset,
        "items": [_row(item, queries) for item in rows],
    }


def load_platform_rows(session: Session, platform: str) -> list[dict[str, Any]]:
    page = list_platform_rows(session, platform, limit=100_000, offset=0)
    return list(page["items"])


def write_platform_workbook(session: Session, path: str | Path) -> dict[str, int]:
    """Write a new workbook. An existing file is replaced and is not reused as input.

    Incomplete cards stay out of the file. A Google row is built only from that
    Google snapshot, so a 2GIS name, phone, address, or category is not copied in.
    """
    book = Workbook()
    counts: dict[str, int] = {}
    first = True
    for platform in PLATFORMS:
        sheet = book.active if first else book.create_sheet()
        first = False
        sheet.title = LABELS[platform]
        rows = [row for row in load_platform_rows(session, platform) if row.get("completeness") == "complete"]
        _write_sheet(sheet, rows)
        counts[platform] = len(rows)
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    temporary = dest.with_name(dest.name + ".partial")
    book.save(temporary)
    os.replace(temporary, dest)
    return counts


def _latest_ids(platform: str):
    """Newest complete card wins. A later incomplete import does not hide it.

    History rows stay in the table. This query only chooses which row to show.
    """
    name_text = func.nullif(func.btrim(CardSnapshot.fields["name"].as_string()), "")
    url_text = func.nullif(func.btrim(CardSnapshot.url), "")
    if platform == "google":
        url_ok = and_(url_text.is_not(None), or_(url_text.ilike("http://%"), url_text.ilike("https://%")))
    else:
        url_ok = url_text.is_not(None)
    complete = case((and_(name_text.is_not(None), url_ok), 0), else_=1)
    rank = func.row_number().over(
        partition_by=CardSnapshot.company_id,
        order_by=(complete.asc(), CardSnapshot.collected_at.desc(), CardSnapshot.id.desc()),
    )
    ranked = (
        select(CardSnapshot.id.label("id"), rank.label("rn"))
        .where(
            CardSnapshot.platform == platform,
            CardSnapshot.company_id.is_not(None),
            CardSnapshot.match_status == "found",
        )
        .subquery()
    )
    return select(ranked.c.id).where(ranked.c.rn == 1)


def _row(snapshot: CardSnapshot, queries: dict[int, str] | None = None) -> dict[str, Any]:
    fields = snapshot.fields if isinstance(snapshot.fields, dict) else {}
    body: dict[str, Any] = {"company_id": snapshot.company_id}
    for key in _TEXT_KEYS:
        body[key] = _text(fields.get(key))
    body["phone"] = _phone(fields.get("phones"))
    body["source_url"] = _text(snapshot.url)
    body["rating"] = _number(fields.get("rating"), whole=False)
    body["reviews_count"] = _number(fields.get("reviews_count"), whole=True)
    body["collected_at"] = snapshot.collected_at.isoformat() if isinstance(snapshot.collected_at, datetime) else None
    verdict = assess_card(snapshot.platform, fields, snapshot.url)
    body["name"] = verdict["name"]
    body["source_url"] = verdict["source_url"]
    body["completeness"] = verdict["completeness"]
    body["completeness_label"] = verdict["completeness_label"]
    body["category"] = _category_text(snapshot, fields)
    body["search_query"] = _search_query(snapshot, queries or {})
    body["source"] = LABELS.get(snapshot.platform, snapshot.platform)
    shown_city, region = _place(snapshot, fields)
    body["city"] = shown_city
    body["region"] = region
    return body


def _write_sheet(sheet: Worksheet, rows: list[dict[str, Any]]) -> None:
    headers = [title for _key, title in COLUMNS]
    keys = [key for key, _title in COLUMNS]
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel(row.get(key)) for key in keys])
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for index, header in enumerate(headers, start=1):
        width = 36 if header == "Ссылка на карточку" else min(40, max(14, len(header) + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width


def _text(value: Any) -> str | None:
    if value is None or isinstance(value, (list, dict, bool)):
        return None
    text = " ".join(str(value).split())
    return text or None


def _phone(value: Any) -> str | None:
    if not isinstance(value, list):
        return None
    numbers = []
    for item in value:
        text = _text(item)
        if text and text not in numbers:
            numbers.append(text)
    return ", ".join(numbers) or None


def _number(value: Any, *, whole: bool) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if whole and isinstance(value, int):
        return value
    if not whole and isinstance(value, (int, float)):
        return value
    return None


def _excel(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    if text[:1] in {"=", "+", "-", "@"}:
        return "'" + text
    return text


def _escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _category_text(snapshot: CardSnapshot, fields: dict[str, Any]) -> str:
    stored = _text(fields.get("category"))
    if stored:
        return stored
    recovered = _recover_single_label(snapshot)
    if recovered:
        return recovered
    return NOT_STATED


def _place(snapshot: CardSnapshot, fields: dict[str, Any]) -> tuple[str, str]:
    """Prefer the city written on the card. The selected search city is a fallback, not an address."""
    sources = snapshot.field_sources if isinstance(snapshot.field_sources, dict) else {}
    marker = sources.get("_collection") if isinstance(sources.get("_collection"), dict) else {}
    verified = _text(fields.get("city"))
    if verified and verified != UNCONFIRMED_CITY:
        return verified, region_of(verified) or NOT_STATED
    selected = _text(marker.get("collection_city"))
    if selected:
        return selected, _text(marker.get("collection_region")) or region_of(selected) or NOT_STATED
    if verified == UNCONFIRMED_CITY:
        return UNCONFIRMED_CITY, NOT_STATED
    return NOT_STATED, NOT_STATED


def _search_query(snapshot: CardSnapshot, queries: dict[int, str]) -> str | None:
    sources = snapshot.field_sources if isinstance(snapshot.field_sources, dict) else {}
    marker = sources.get("_collection")
    if not isinstance(marker, dict):
        return None
    direct = _text(marker.get("search_query"))
    if direct:
        return direct
    batch_id = _batch_id(marker.get("key"))
    if batch_id is None:
        return None
    return _text(queries.get(batch_id))


def _batch_queries(session: Session, snapshots: list[CardSnapshot]) -> dict[int, str]:
    needed: set[int] = set()
    for snapshot in snapshots:
        sources = snapshot.field_sources if isinstance(snapshot.field_sources, dict) else {}
        marker = sources.get("_collection")
        if not isinstance(marker, dict) or _text(marker.get("search_query")):
            continue
        batch_id = _batch_id(marker.get("key"))
        if batch_id is not None:
            needed.add(batch_id)
    if not needed:
        return {}
    found = session.scalars(select(CollectionBatch).where(CollectionBatch.id.in_(needed))).all()
    queries: dict[int, str] = {}
    for row in found:
        text = _text(row.category)
        if text:
            queries[row.id] = text
    return queries


def _batch_id(key: Any) -> int | None:
    if not isinstance(key, str):
        return None
    parts = key.split(":")
    if len(parts) < 2 or parts[0] not in _QUERY_PREFIXES or not parts[1].isdigit():
        return None
    return int(parts[1])


def _recover_single_label(snapshot: CardSnapshot) -> str | None:
    """Read one rubric from a saved page when the collector recorded unknown_label.

    Several labels stay unused. Nothing is written back to the snapshot.
    """
    sources = snapshot.field_sources if isinstance(snapshot.field_sources, dict) else {}
    source = sources.get("category")
    if not isinstance(source, dict) or source.get("reason") != "unknown_label":
        return None
    marker = sources.get("_collection")
    raw = marker.get("raw_path") if isinstance(marker, dict) else None
    if not isinstance(raw, str) or not raw.strip():
        return None
    path = Path(raw)
    if not path.is_absolute():
        path = _REPO / path
    try:
        resolved = path.resolve()
        resolved.relative_to(_REPO.resolve())
    except (OSError, ValueError):
        return None
    if not resolved.is_file():
        return None
    try:
        html = resolved.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None
    labels = _labels_from_html(snapshot.platform, html)
    if len(labels) != 1:
        return None
    return labels[0]


def _labels_from_html(platform: str, html: str) -> list[str]:
    parser = _CategoryParser(platform)
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return []
    return parser.labels


class _CategoryParser(HTMLParser):
    def __init__(self, platform: str) -> None:
        super().__init__(convert_charrefs=True)
        self.platform = platform
        self.labels: list[str] = []
        self._tag: str | None = None
        self._depth = 0
        self._parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = {key: value or "" for key, value in attrs}
        classes = set(values.get("class", "").split())
        href = values.get("href", "")
        start = False
        if self.platform == "2gis" and tag == "a" and "rubricId/" in href:
            start = True
        elif self.platform == "google" and tag == "button" and "DkEaL" in classes:
            start = True
        elif self.platform == "yandex" and "orgpage-categories-info-view" in classes:
            start = True
        elif self.platform == "yandex" and tag == "a" and "business-categories-view__category" in classes:
            start = True
        if self._tag is None and start:
            self._tag = tag
            self._depth = 1
            self._parts = []
        elif self._tag == tag:
            self._depth += 1

    def handle_endtag(self, tag: str) -> None:
        if self._tag != tag:
            return
        self._depth -= 1
        if self._depth > 0:
            return
        text = " ".join("".join(self._parts).split())
        if text and text not in self.labels:
            self.labels.append(text)
        self._tag = None
        self._parts = []

    def handle_data(self, data: str) -> None:
        if self._tag is not None:
            self._parts.append(data)
